from __future__ import annotations

import csv
import re
import unicodedata
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from statistics import median
from typing import Any

from .models import DataQuality, LoadedData, MeasurementKind, Reading
from .units import convert_energy_to_kwh, convert_power_to_kw


ALIASES = {
    "date": {
        "date", "datetime", "timestamp", "horodatage", "jour", "periode",
        "date_heure", "dateheure", "time",
    },
    "energy": {
        "energie", "energie_wh", "energie_kwh", "energie_mwh",
        "energy", "energy_wh", "energy_kwh", "energy_mwh", "wh", "kwh", "mwh",
        "consommation", "consommation_kwh", "conso", "conso_kwh",
        "consommation_wh", "consommation_mwh", "conso_wh", "conso_mwh",
        "index", "index_wh", "index_kwh", "index_mwh",
        "releve", "releve_wh", "releve_kwh", "releve_mwh",
    },
    "production": {
        "production", "quantite_produite", "quantite", "output",
        "volume_production", "unites_produites", "production_unites",
    },
    "production_active": {
        "production_active", "production_en_cours", "active", "is_producing",
    },
    "temperature": {
        "outside_temperature", "outside_temperature_c", "temperature_exterieure",
        "temperature_exterieure_c", "temperature", "temperature_c", "temp_c",
    },
    "shift": {"shift", "equipe", "poste"},
    "product_type": {"product_type", "type_produit", "produit"},
    "tariff": {
        "tarif", "tarif_kwh", "prix_kwh", "cout_kwh", "cost_per_kwh",
        "tariff", "tariff_per_kwh", "eur_kwh",
    },
    "power": {
        "puissance", "puissance_w", "puissance_kw", "puissance_mw",
        "power", "power_w", "power_kw", "power_mw", "w", "kw", "mw",
        "demande_w", "demande_kw", "demande_mw",
        "demand_w", "demand_kw", "demand_mw",
    },
}

DATE_FORMATS = (
    "%Y-%m-%dT%H:%M:%S",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%dT%H:%M",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%d",
    "%d/%m/%Y %H:%M:%S",
    "%d/%m/%Y %H:%M",
    "%d/%m/%Y",
    "%d-%m-%Y %H:%M:%S",
    "%d-%m-%Y %H:%M",
    "%d-%m-%Y",
)


class DataError(ValueError):
    """Erreur de donnees destinee a etre affichee par la CLI."""


def normalize_header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def parse_number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("\u00a0", "").replace(" ", "")
    if not text or text.lower() in {"na", "n/a", "nan", "null", "none", "-"}:
        return None
    text = re.sub(r"[^0-9,\.\-+]", "", text)
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def parse_date(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        if value.tzinfo is None:
            return value
        return value.astimezone(timezone.utc).replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = str(value or "").strip()
    if not text:
        return None
    iso_candidate = text[:-1] + "+00:00" if text.endswith("Z") else text
    try:
        parsed = datetime.fromisoformat(iso_candidate)
        if parsed.tzinfo is None:
            return parsed
        return parsed.astimezone(timezone.utc).replace(tzinfo=None)
    except ValueError:
        pass
    for date_format in DATE_FORMATS:
        try:
            return datetime.strptime(text, date_format)
        except ValueError:
            continue
    return None


def parse_bool(value: Any) -> bool | None:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in {0, 1}:
        return bool(value)
    text = normalize_header(value)
    if text in {"1", "true", "yes", "oui", "active", "actif"}:
        return True
    if text in {"0", "false", "no", "non", "inactive", "inactif"}:
        return False
    return None


def _is_missing_value(value: Any) -> bool:
    if value is None:
        return True
    return str(value).strip().lower() in {"", "na", "n/a", "nan", "null", "none", "-"}


def _timestamp_has_timezone(value: Any) -> bool:
    if isinstance(value, datetime):
        return value.tzinfo is not None and value.utcoffset() is not None
    text = str(value or "").strip()
    return text.endswith("Z") or bool(re.search(r"[+-]\d{2}:?\d{2}$", text))


def _record_optional_quality(
    quality: DataQuality, *, name: str, raw: Any, parsed: Any
) -> None:
    target = (
        quality.missing_values_by_column
        if _is_missing_value(raw)
        else quality.invalid_values_by_column
        if parsed is None
        else None
    )
    if target is not None:
        target[name] = target.get(name, 0) + 1


def _read_csv(path: Path) -> tuple[list[str], list[list[Any]]]:
    last_error: UnicodeDecodeError | None = None
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                sample = handle.read(8192)
                handle.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                rows = list(csv.reader(handle, dialect))
            break
        except UnicodeDecodeError as exc:
            last_error = exc
    else:
        raise DataError(f"Encodage CSV non reconnu: {last_error}")
    if not rows:
        raise DataError("Le fichier est vide.")
    return [str(item).strip() for item in rows[0]], rows[1:]


def _read_xlsx(path: Path) -> tuple[list[str], list[list[Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DataError(
            "La lecture XLSX necessite openpyxl. Installez-la avec: "
            "python -m pip install -r requirements.txt"
        ) from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    try:
        header = next(iterator)
    except StopIteration as exc:
        workbook.close()
        raise DataError("Le classeur est vide.") from exc
    rows = [list(row) for row in iterator]
    workbook.close()
    return [str(item).strip() if item is not None else "" for item in header], rows


def _resolve_column(
    headers: list[str], kind: str, explicit: str | None, required: bool = False
) -> int | None:
    normalized = [normalize_header(header) for header in headers]
    if explicit:
        wanted = normalize_header(explicit)
        if wanted not in normalized:
            raise DataError(
                f"Colonne '{explicit}' introuvable. Colonnes disponibles: "
                + ", ".join(headers)
            )
        return normalized.index(wanted)
    for alias in ALIASES[kind]:
        if alias in normalized:
            return normalized.index(alias)
    if required:
        raise DataError(
            f"Colonne {kind} introuvable. Colonnes disponibles: "
            + ", ".join(headers)
            + f". Utilisez --{kind}-column pour la preciser."
        )
    return None


def _value_at(row: list[Any], index: int | None) -> Any:
    return row[index] if index is not None and index < len(row) else None


def _detect_energy_mode(header: str, requested: str) -> str:
    if requested != "auto":
        return requested
    normalized = normalize_header(header)
    if normalized.startswith(("index", "releve")):
        return "cumulative"
    return "interval"


def _resolve_unit(header: str, explicit: str | None, *, quantity: str) -> str:
    normalized = normalize_header(header)
    supported = ("wh", "kwh", "mwh") if quantity == "energy" else ("w", "kw", "mw")
    if explicit is not None:
        unit = explicit.lower()
        if unit not in supported:
            raise DataError(f"Unite {quantity} non prise en charge: {explicit}.")
        return unit
    tokens = set(normalized.split("_"))
    matches = [unit for unit in supported if unit in tokens or normalized == unit]
    if len(matches) == 1:
        return matches[0]
    option = "--energy-unit" if quantity == "energy" else "--power-unit"
    raise DataError(
        f"Unite absente ou ambigue pour la colonne '{header}'. "
        f"Precisez-la avec {option}."
    )


def _reading_signature(reading: Reading) -> tuple[float | None, ...]:
    return (
        reading.energy_kwh,
        reading.production,
        reading.production_active,
        reading.outside_temperature_c,
        reading.shift,
        reading.product_type,
        reading.tariff_per_kwh,
        reading.power_kw,
    )


def _deduplicate(readings: list[Reading], quality: DataQuality) -> list[Reading]:
    unique: list[Reading] = []
    by_timestamp: dict[datetime, Reading] = {}
    for reading in readings:
        previous = by_timestamp.get(reading.timestamp)
        if previous is None:
            by_timestamp[reading.timestamp] = reading
            unique.append(reading)
            continue
        if _reading_signature(previous) == _reading_signature(reading):
            quality.identical_duplicates_removed += 1
            continue
        quality.conflicting_duplicates += 1
        raise DataError(
            "Doublon conflictuel au timestamp "
            f"{reading.timestamp.isoformat()} (lignes {previous.source_row} et "
            f"{reading.source_row}). Corrigez ou dedupliquez explicitement la source."
        )
    if quality.identical_duplicates_removed:
        count = quality.identical_duplicates_removed
        quality.processing_log.append(
            f"{count} doublon(s) strictement identique(s) supprime(s)."
        )
    return unique


def _populate_time_quality(readings: list[Reading], quality: DataQuality) -> list[float]:
    deltas_seconds = [
        (current.timestamp - previous.timestamp).total_seconds()
        for previous, current in zip(readings, readings[1:])
    ]
    positive_deltas = [value for value in deltas_seconds if value > 0]
    if not positive_deltas:
        return []
    looks_monthly = all(27 * 86400 <= value <= 32 * 86400 for value in positive_deltas)
    rounded_deltas = [round(value, 3) for value in positive_deltas]
    counts = Counter(rounded_deltas)
    typical = (
        median(positive_deltas)
        if looks_monthly
        else min(counts, key=lambda value: (-counts[value], value))
    )
    tolerance = 4 * 86400 if looks_monthly else max(1.0, typical * 0.01)
    quality.inferred_frequency_minutes = typical / 60.0
    if looks_monthly:
        quality.detected_frequencies_minutes = [typical / 60.0]
    else:
        runs: list[tuple[float, int]] = []
        for value in rounded_deltas:
            if runs and abs(runs[-1][0] - value) <= tolerance:
                runs[-1] = (runs[-1][0], runs[-1][1] + 1)
            else:
                runs.append((value, 1))
        sustained = [run for run in runs if run[1] >= 2]
        recurring = sorted({value for value, _ in sustained})
        quality.detected_frequencies_minutes = [value / 60.0 for value in recurring]
        previous_frequency: float | None = None
        for value, _ in sustained:
            if previous_frequency is not None and abs(value - previous_frequency) > tolerance:
                quality.frequency_change_count += 1
            previous_frequency = value
    quality.irregular_interval_count = sum(
        abs(value - typical) > tolerance for value in positive_deltas
    )
    recurring_seconds = [value * 60 for value in quality.detected_frequencies_minutes]
    quality.gap_count = sum(
        value > typical * 1.5
        and not any(abs(value - recurring) <= tolerance for recurring in recurring_seconds)
        for value in positive_deltas
    )
    span = (readings[-1].timestamp - readings[0].timestamp).total_seconds()
    expected_rows = round(span / typical) + 1
    if expected_rows > 0 and quality.frequency_change_count == 0:
        quality.coverage_ratio = min(1.0, len(readings) / expected_rows)
    quality.processing_log.append(
        f"Frequence nominale detectee: {quality.inferred_frequency_minutes:g} minute(s)."
    )
    if quality.frequency_change_count:
        values = ", ".join(f"{value:g}" for value in quality.detected_frequencies_minutes)
        quality.processing_log.append(
            f"{quality.frequency_change_count} changement(s) durable(s) de frequence "
            f"detecte(s) ({values} minute(s)); couverture globale non estimee."
        )
    return positive_deltas


def _assign_interval_hours(
    readings: list[Reading], *, interval_minutes: float | None, quality: DataQuality
) -> None:
    if interval_minutes is not None:
        if interval_minutes <= 0:
            raise DataError("--interval-minutes doit etre strictement positif.")
        hours = interval_minutes / 60.0
        for reading in readings:
            reading.interval_hours = hours
        return
    if len(readings) < 2:
        return
    if quality.frequency_change_count:
        quality.processing_log.append(
            "Duree nominale non attribuee: plusieurs frequences durables sont presentes."
        )
        return
    deltas = [
        (current.timestamp - previous.timestamp).total_seconds() / 3600.0
        for previous, current in zip(readings, readings[1:])
    ]
    typical = median(deltas)
    # Une energie par intervalle reste attachee a l'intervalle nominal : un trou
    # represente des lignes absentes, pas une mesure couvrant silencieusement tout le trou.
    for reading in readings:
        reading.interval_hours = typical


def _coverage_bounds(
    readings: list[Reading], *, timestamp_position: str, cumulative: bool
) -> tuple[datetime | None, datetime | None]:
    if not readings:
        return None, None
    if cumulative:
        first_hours = readings[0].interval_hours
        start = (
            readings[0].timestamp
            if first_hours is None
            else readings[0].timestamp - timedelta(hours=first_hours)
        )
        return start, readings[-1].timestamp
    first_hours = readings[0].interval_hours
    last_hours = readings[-1].interval_hours
    if first_hours is None or last_hours is None:
        return None, None
    if timestamp_position == "start":
        return readings[0].timestamp, readings[-1].timestamp + timedelta(hours=last_hours)
    return readings[0].timestamp - timedelta(hours=first_hours), readings[-1].timestamp


def load_data(
    path: str | Path,
    *,
    date_column: str | None = None,
    energy_column: str | None = None,
    production_column: str | None = None,
    production_active_column: str | None = None,
    temperature_column: str | None = None,
    shift_column: str | None = None,
    product_type_column: str | None = None,
    tariff_column: str | None = None,
    power_column: str | None = None,
    energy_mode: str = "auto",
    interval_minutes: float | None = None,
    energy_unit: str | None = None,
    power_unit: str | None = None,
    timestamp_position: str = "auto",
) -> LoadedData:
    source = Path(path).expanduser()
    if not source.is_file():
        raise DataError(f"Fichier introuvable: {source}")
    if timestamp_position not in {"auto", "start", "end"}:
        raise DataError("timestamp_position doit valoir auto, start ou end.")
    suffix = source.suffix.lower()
    if suffix == ".csv":
        headers, rows = _read_csv(source)
    elif suffix == ".xlsx":
        headers, rows = _read_xlsx(source)
    else:
        raise DataError("Format non pris en charge. Utilisez un fichier .csv ou .xlsx.")

    indices = {
        "date": _resolve_column(headers, "date", date_column, required=True),
        "energy": _resolve_column(headers, "energy", energy_column),
        "production": _resolve_column(headers, "production", production_column),
        "production_active": _resolve_column(
            headers, "production_active", production_active_column
        ),
        "temperature": _resolve_column(headers, "temperature", temperature_column),
        "shift": _resolve_column(headers, "shift", shift_column),
        "product_type": _resolve_column(headers, "product_type", product_type_column),
        "tariff": _resolve_column(headers, "tariff", tariff_column),
        "power": _resolve_column(headers, "power", power_column),
    }
    if indices["energy"] is None and indices["power"] is None:
        raise DataError(
            "Aucune mesure energetique reconnue. Fournissez une colonne kWh "
            "(--energy-column) ou kW (--power-column)."
        )

    if indices["energy"] is not None:
        resolved_energy_unit = _resolve_unit(
            headers[indices["energy"]], energy_unit, quantity="energy"
        )
        selected_mode = _detect_energy_mode(headers[indices["energy"]], energy_mode)
        measurement_kind = (
            MeasurementKind.CUMULATIVE_ENERGY
            if selected_mode == "cumulative"
            else MeasurementKind.ENERGY_PER_INTERVAL
        )
    else:
        resolved_energy_unit = None
        if energy_mode not in {"auto", "interval"}:
            raise DataError(
                "--energy-mode cumulative exige une colonne d'index energetique."
            )
        selected_mode = "power"
        measurement_kind = MeasurementKind.POWER

    if (
        measurement_kind is MeasurementKind.CUMULATIVE_ENERGY
        and timestamp_position == "start"
    ):
        raise DataError(
            "Un index cumulatif est un releve instantane de fin d'intervalle; "
            "--timestamp-position start est incompatible."
        )

    resolved_timestamp_position = (
        "end"
        if measurement_kind is MeasurementKind.CUMULATIVE_ENERGY
        else "start"
        if timestamp_position == "auto"
        else timestamp_position
    )

    resolved_power_unit = (
        _resolve_unit(headers[indices["power"]], power_unit, quantity="power")
        if indices["power"] is not None
        else None
    )

    warnings: list[str] = []
    quality = DataQuality()
    parsed: list[Reading] = []
    discarded = 0
    for row_number, row in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in row):
            quality.blank_rows += 1
            continue
        raw_timestamp = _value_at(row, indices["date"])
        timestamp = parse_date(raw_timestamp)
        measurement_index = indices["energy"] if indices["energy"] is not None else indices["power"]
        measurement = parse_number(_value_at(row, measurement_index))
        if timestamp is None:
            quality.invalid_timestamp_rows += 1
            discarded += 1
            continue
        if _timestamp_has_timezone(raw_timestamp):
            quality.timezone_normalized_rows += 1
        if measurement is None:
            quality.missing_measurement_rows += 1
            discarded += 1
            continue
        raw_production = _value_at(row, indices["production"])
        production = parse_number(raw_production)
        if indices["production"] is not None:
            _record_optional_quality(
                quality, name="production", raw=raw_production, parsed=production
            )
        raw_active = _value_at(row, indices["production_active"])
        production_active = parse_bool(raw_active)
        if indices["production_active"] is not None:
            _record_optional_quality(
                quality,
                name="production_active",
                raw=raw_active,
                parsed=production_active,
            )
        raw_temperature = _value_at(row, indices["temperature"])
        outside_temperature = parse_number(raw_temperature)
        if indices["temperature"] is not None:
            _record_optional_quality(
                quality,
                name="outside_temperature_c",
                raw=raw_temperature,
                parsed=outside_temperature,
            )
            if outside_temperature is not None and not -100 <= outside_temperature <= 80:
                quality.invalid_values_by_column["outside_temperature_c"] = (
                    quality.invalid_values_by_column.get("outside_temperature_c", 0) + 1
                )
                outside_temperature = None
        shift_value = _value_at(row, indices["shift"])
        shift = str(shift_value).strip() if shift_value not in (None, "") else None
        if indices["shift"] is not None:
            _record_optional_quality(quality, name="shift", raw=shift_value, parsed=shift)
        product_value = _value_at(row, indices["product_type"])
        product_type = (
            str(product_value).strip() if product_value not in (None, "") else None
        )
        if indices["product_type"] is not None:
            _record_optional_quality(
                quality, name="product_type", raw=product_value, parsed=product_type
            )
        raw_tariff = _value_at(row, indices["tariff"])
        tariff = parse_number(raw_tariff)
        if indices["tariff"] is not None:
            _record_optional_quality(
                quality, name="tariff_per_kwh", raw=raw_tariff, parsed=tariff
            )
        raw_power_value = _value_at(row, indices["power"])
        raw_power = parse_number(raw_power_value)
        if indices["power"] is not None and indices["energy"] is not None:
            _record_optional_quality(
                quality, name="power", raw=raw_power_value, parsed=raw_power
            )
        power = (
            convert_power_to_kw(raw_power, resolved_power_unit)
            if raw_power is not None and resolved_power_unit is not None
            else None
        )
        energy = (
            convert_energy_to_kwh(measurement, resolved_energy_unit)
            if indices["energy"] is not None and resolved_energy_unit is not None
            else 0.0
        )
        if production is not None and production < 0:
            quality.invalid_values_by_column["production"] = (
                quality.invalid_values_by_column.get("production", 0) + 1
            )
            production = None
        if tariff is not None and tariff < 0:
            quality.invalid_values_by_column["tariff_per_kwh"] = (
                quality.invalid_values_by_column.get("tariff_per_kwh", 0) + 1
            )
            tariff = None
        primary_power_is_impossible = (
            indices["energy"] is None and power is not None and power < 0
        )
        if indices["energy"] is not None and power is not None and power < 0:
            quality.invalid_values_by_column["power"] = (
                quality.invalid_values_by_column.get("power", 0) + 1
            )
            power = None
        if energy < 0 or primary_power_is_impossible:
            quality.impossible_value_rows += 1
            discarded += 1
            continue
        parsed.append(
            Reading(
                timestamp=timestamp,
                energy_kwh=energy,
                production=production,
                production_active=production_active,
                outside_temperature_c=outside_temperature,
                shift=shift,
                product_type=product_type,
                tariff_per_kwh=tariff,
                power_kw=power,
                source_row=row_number,
            )
        )
    if not parsed:
        raise DataError("Aucune ligne exploitable: verifiez les dates et les mesures.")
    if quality.blank_rows:
        discarded += quality.blank_rows
        quality.processing_log.append(
            f"{quality.blank_rows} ligne(s) vide(s) ignoree(s)."
        )
    quality.out_of_order_rows = sum(
        current.timestamp < previous.timestamp
        for previous, current in zip(parsed, parsed[1:])
    )
    if quality.out_of_order_rows:
        quality.processing_log.append(
            f"{quality.out_of_order_rows} rupture(s) d'ordre chronologique detectee(s); "
            "les lignes ont ete triees par timestamp."
        )
    parsed.sort(key=lambda reading: reading.timestamp)
    parsed = _deduplicate(parsed, quality)
    discarded += quality.identical_duplicates_removed
    _populate_time_quality(parsed, quality)
    _assign_interval_hours(
        parsed, interval_minutes=interval_minutes, quality=quality
    )
    if resolved_energy_unit is not None and resolved_energy_unit != "kwh":
        quality.processing_log.append(
            f"Energie convertie de {resolved_energy_unit} vers kWh."
        )
    if resolved_power_unit is not None and resolved_power_unit != "kw":
        quality.processing_log.append(
            f"Puissance convertie de {resolved_power_unit} vers kW."
        )

    if selected_mode == "cumulative":
        converted: list[Reading] = []
        resets = 0
        for previous, current in zip(parsed, parsed[1:]):
            delta = current.energy_kwh - previous.energy_kwh
            if delta < 0:
                resets += 1
                continue
            converted.append(
                Reading(
                    timestamp=current.timestamp,
                    energy_kwh=delta,
                    production=current.production,
                    production_active=current.production_active,
                    outside_temperature_c=current.outside_temperature_c,
                    shift=current.shift,
                    product_type=current.product_type,
                    tariff_per_kwh=current.tariff_per_kwh,
                    power_kw=current.power_kw,
                    source_row=current.source_row,
                    interval_hours=(
                        current.timestamp - previous.timestamp
                    ).total_seconds() / 3600.0,
                )
            )
        if not converted:
            raise DataError("Impossible de calculer les ecarts de l'index cumulatif.")
        parsed = converted
        discarded += 1 + resets
        quality.processing_log.append(
            "La premiere valeur de l'index sert de reference et n'est pas une energie d'intervalle."
        )
        if resets:
            warnings.append(f"{resets} remise(s) a zero d'index ignoree(s).")
    elif selected_mode == "power":
        if interval_minutes is None:
            if len(parsed) < 2:
                raise DataError(
                    "Une mesure de puissance isolee exige --interval-minutes pour calculer les kWh."
                )
            if quality.irregular_interval_count:
                raise DataError(
                    "Les timestamps de puissance sont irreguliers. Precisez explicitement "
                    "--interval-minutes ou fournissez une colonne d'energie par intervalle."
                )
        for reading in parsed:
            if reading.interval_hours is None or reading.power_kw is None:
                raise DataError("Duree d'intervalle indisponible pour convertir les kW en kWh.")
            reading.energy_kwh = reading.power_kw * reading.interval_hours
        quality.processing_log.append(
            "Puissance convertie en energie par la formule kWh = kW x duree en heures."
        )

    coverage_start, coverage_end = _coverage_bounds(
        parsed,
        timestamp_position=resolved_timestamp_position,
        cumulative=measurement_kind is MeasurementKind.CUMULATIVE_ENERGY,
    )
    coverage_bounds_method = (
        "cumulative_reading_differences"
        if measurement_kind is MeasurementKind.CUMULATIVE_ENERGY
        else "explicit_interval"
        if interval_minutes is not None
        else "inferred_nominal_interval"
        if coverage_start is not None and coverage_end is not None
        else "unavailable"
    )
    if measurement_kind is MeasurementKind.CUMULATIVE_ENERGY:
        quality.processing_log.append(
            "Les timestamps d'index sont traites comme fins d'intervalles entre deux releves."
        )
    elif timestamp_position == "auto":
        quality.processing_log.append(
            "Les timestamps sont traites par defaut comme debuts d'intervalles; "
            "utiliser --timestamp-position end si la source suit l'autre convention."
        )

    if quality.identical_duplicates_removed:
        warnings.append(
            f"{quality.identical_duplicates_removed} doublon(s) strictement identique(s) supprime(s)."
        )
    if quality.irregular_interval_count:
        warnings.append(
            f"{quality.irregular_interval_count} intervalle(s) different de la frequence nominale."
        )
    if quality.out_of_order_rows:
        warnings.append(
            f"{quality.out_of_order_rows} rupture(s) d'ordre source corrigee(s) et tracee(s)."
        )
    if quality.frequency_change_count:
        warnings.append(
            f"{quality.frequency_change_count} changement(s) durable(s) de frequence detecte(s); "
            "les analyses dependantes d'une duree uniforme sont limitees."
        )
    if quality.gap_count:
        warnings.append(f"{quality.gap_count} trou(s) temporel(s) probable(s) detecte(s).")
    if quality.impossible_value_rows:
        quality.processing_log.append(
            f"{quality.impossible_value_rows} ligne(s) avec valeur impossible ecartee(s)."
        )
    if quality.invalid_timestamp_rows:
        quality.processing_log.append(
            f"{quality.invalid_timestamp_rows} ligne(s) avec timestamp invalide ecartee(s)."
        )
    if quality.missing_measurement_rows:
        quality.processing_log.append(
            f"{quality.missing_measurement_rows} ligne(s) sans mesure energetique ecartee(s)."
        )
    if quality.timezone_normalized_rows:
        quality.processing_log.append(
            f"{quality.timezone_normalized_rows} timestamp(s) avec fuseau converti(s) en UTC."
        )
    for column, count in sorted(quality.missing_values_by_column.items()):
        quality.processing_log.append(
            f"{count} valeur(s) contextuelle(s) manquante(s) dans {column}."
        )
    for column, count in sorted(quality.invalid_values_by_column.items()):
        quality.processing_log.append(
            f"{count} valeur(s) contextuelle(s) invalide(s) dans {column}."
        )
    if discarded:
        warnings.append(
            f"{discarded} ligne(s) ecartee(s) de l'analyse "
            "(vides, invalides, incompletes, references ou dedupliquees)."
        )
    columns = {
        kind: headers[index]
        for kind, index in indices.items()
        if index is not None
    }
    return LoadedData(
        readings=parsed,
        input_rows=len(rows),
        discarded_rows=discarded,
        warnings=warnings,
        columns=columns,
        source_units={
            kind: unit
            for kind, unit in {
                "energy": resolved_energy_unit,
                "power": resolved_power_unit,
            }.items()
            if unit is not None
        },
        energy_mode=selected_mode,
        measurement_kind=measurement_kind,
        quality=quality,
        timestamp_position=resolved_timestamp_position,
        coverage_start=coverage_start,
        coverage_end=coverage_end,
        coverage_bounds_method=coverage_bounds_method,
    )

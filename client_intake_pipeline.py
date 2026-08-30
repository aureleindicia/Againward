"""Intake deterministic et traçable d'un dépôt hétérogène client.

Ce module ne produit ni diagnostic, ni cause, ni recommandation.  Il prépare un
cas canonique compact pour que Codex conduise l'investigation, en conservant les
sources et les décisions de normalisation nécessaires à l'audit.
"""
from __future__ import annotations

import csv
import hashlib
import json
import mimetypes
import re
import shutil
import unicodedata
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from statistics import median
from typing import Any, Iterable

from energy_mvp.io import ALIASES, parse_date, parse_number


CASE_DIRECTORIES = (
    "raw",
    "normalized",
    "derived",
    "evidence",
    "investigation",
    "outputs",
    "logs",
)
ROLE_VALUES = {
    "ENERGY_INTERVAL_SERIES", "ENERGY_MONTHLY", "POWER_SERIES", "PRODUCTION",
    "SCHEDULE", "TARIFF", "INVOICE", "WEATHER", "MAINTENANCE",
    "MACHINE_METADATA", "PROCESS_CONTEXT", "TEXT_NOTE", "UNKNOWN", "IRRELEVANT",
}
USABILITY_VALUES = {"usable", "partially_usable", "irrelevant", "unreadable"}
TRANSFORMATION_STATUS = {"AUTO_FIX_SAFE", "FLAG_ONLY", "MATERIAL_AMBIGUITY", "UNUSABLE"}
REQUEST_TYPES = {
    "INFER_AUTOMATICALLY", "ASK_CLIENT", "REQUEST_EXISTING_DOCUMENT",
    "REQUEST_TECHNICAL_EVIDENCE", "OPTIONAL_FUTURE_INSTRUMENTATION",
}
REQUEST_IMPORTANCE = {"BLOCKING", "NON_BLOCKING", "OPTIONAL"}
MAX_DEFAULT_QUESTIONS = 3

_TEXT_EXTENSIONS = {".txt", ".md", ".json", ".log", ".pdf"}
_SPREADSHEET_EXTENSIONS = {".csv", ".xlsx", ".xls"}
_IRRELEVANT_NAMES = {"photo", "image", "cv", "logo", "brochure", "catalogue"}
_ROLE_KEYWORDS = {
    "SCHEDULE": ("schedule", "planning", "horaire", "ouverture", "shift"),
    "TARIFF": ("tarif", "tariff", "prix", "price", "contract"),
    "INVOICE": ("facture", "invoice", "bill"),
    "WEATHER": ("weather", "meteo", "météo", "temperature", "temp"),
    "MAINTENANCE": ("maintenance", "intervention", "frigoriste", "repair"),
    "MACHINE_METADATA": ("machine", "equipement", "équipement", "asset", "parc"),
    "PROCESS_CONTEXT": ("process", "procédé", "production", "activité", "activite"),
}


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _slug(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    normalized = "".join(char for char in normalized if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "_", normalized.lower()).strip("_")


def _safe_relative_files(source: Path) -> Iterable[Path]:
    for path in sorted(source.rglob("*")):
        if not path.is_file() or path.is_symlink():
            continue
        relative = path.relative_to(source)
        if any(part in {".", ".."} for part in relative.parts):
            continue
        yield path


def create_client_case(case_id: str, *, root: str | Path = "client_cases") -> dict[str, Any]:
    """Crée un dossier client isolé; il refuse tout écrasement."""

    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,63}", case_id):
        raise ValueError("case_id doit contenir 1 à 64 caractères ASCII sûrs.")
    target = Path(root) / case_id
    if target.exists():
        raise FileExistsError(f"Le cas existe déjà et ne sera pas modifié: {target}")
    target.mkdir(parents=True)
    for directory in CASE_DIRECTORIES:
        (target / directory).mkdir()
    manifest = {
        "schema_version": 1,
        "case_id": case_id,
        "created_at_utc": _utc_now(),
        "status": "awaiting_raw_drop",
        "retention": {"policy": "not_configured", "review_required_before_deletion": True},
        "security": {
            "raw_inputs_immutable_after_ingestion": True,
            "raw_inputs_must_not_be_committed": True,
            "secrets_allowed_in_case_artifacts": False,
        },
        "architecture": {
            "codex": "choisit l'investigation, les hypothèses, les demandes et les décisions",
            "python": "extrait, normalise, mesure, valide et conserve la provenance",
            "deterministic_diagnosis": False,
        },
    }
    _write_json(target / "case_manifest.json", manifest)
    _write_json(target / "investigation" / "case_state.json", {
        "schema_version": 1,
        "status": "awaiting_intake",
        "findings": [],
        "hypotheses": [],
        "client_answers": [],
        "unresolved_requests": [],
        "history": [],
        "ground_truth_used": False,
    })
    return manifest


def _read_csv_rows(path: Path) -> list[list[Any]]:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                sample = handle.read(8192)
                handle.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                return [list(row) for row in csv.reader(handle, dialect)]
        except UnicodeDecodeError:
            continue
    raise ValueError("Encodage texte non pris en charge.")


def _read_xlsx_sheets(path: Path) -> dict[str, list[list[Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - requirements guard
        raise ValueError("openpyxl est requis pour les fichiers XLSX.") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return {
            worksheet.title: [list(row) for row in worksheet.iter_rows(values_only=True)]
            for worksheet in workbook.worksheets
        }
    finally:
        workbook.close()


def _header_score(row: list[Any]) -> tuple[int, dict[str, int]]:
    headers = [_slug(str(value or "")) for value in row]
    found: dict[str, int] = {}
    for semantic, aliases in ALIASES.items():
        for index, header in enumerate(headers):
            if header in aliases:
                found[semantic] = index
                break
    score = 0
    if "date" in found:
        score += 4
    if "energy" in found or "power" in found:
        score += 4
    score += sum(1 for semantic in ("production", "temperature", "tariff") if semantic in found)
    return score, found


def _candidate_table(rows: list[list[Any]], *, sheet: str | None) -> dict[str, Any] | None:
    best: tuple[int, int, dict[str, int]] | None = None
    for index, row in enumerate(rows[:80]):
        score, fields = _header_score(row)
        if best is None or score > best[0]:
            best = (score, index, fields)
    if best is None or best[0] < 4 or "date" not in best[2] or not ({"energy", "power", "production"} & set(best[2])):
        return None
    _, header_row, fields = best
    headers = [str(value or "").strip() for value in rows[header_row]]
    data_rows = [row for row in rows[header_row + 1:] if any(str(value or "").strip() for value in row)]
    if not data_rows:
        return None
    return {
        "sheet": sheet,
        "header_row": header_row + 1,
        "headers": headers,
        "field_indices": fields,
        "rows": data_rows,
    }


def _role_from_table(table: dict[str, Any]) -> tuple[str, list[str]]:
    fields = table["field_indices"]
    if "power" in fields:
        return "POWER_SERIES", ["recognized power header", "recognized timestamp header"]
    if "energy" in fields:
        timestamps = [parse_date(row[fields["date"]] if fields["date"] < len(row) else None) for row in table["rows"]]
        dates = sorted(item for item in timestamps if item is not None)
        if len(dates) >= 3:
            intervals = [(right - left).total_seconds() / 86400 for left, right in zip(dates, dates[1:])]
            if intervals and median(intervals) >= 27:
                return "ENERGY_MONTHLY", ["recognized energy header", "monthly timestamp spacing"]
        return "ENERGY_INTERVAL_SERIES", ["recognized energy header", "recognized timestamp header"]
    if "production" in fields:
        return "PRODUCTION", ["recognized production header", "recognized timestamp header"]
    return "UNKNOWN", []


def _role_from_name(path: Path) -> tuple[str, list[str]]:
    stem = _slug(path.stem)
    if any(word in stem for word in _IRRELEVANT_NAMES):
        return "IRRELEVANT", ["filename indicates non-operational material"]
    for role, keywords in _ROLE_KEYWORDS.items():
        if any(_slug(keyword) in stem for keyword in keywords):
            return role, ["filename keyword"]
    if path.suffix.lower() in _TEXT_EXTENSIONS:
        return "TEXT_NOTE", ["readable text-like file without stronger evidence"]
    return "UNKNOWN", []


def _inferred_unit(table: dict[str, Any], *, role: str) -> dict[str, Any]:
    fields = table["field_indices"]
    measure_key = "power" if "power" in fields else "energy"
    header = _slug(table["headers"][fields[measure_key]])
    explicit = next((unit for unit in ("mwh", "kwh", "wh", "mw", "kw", "w") if unit in header.split("_")), None)
    if explicit:
        return {"status": "AUTO_FIX_SAFE", "unit": explicit, "measurement_kind": "power" if measure_key == "power" else "energy_per_interval", "basis": "explicit unit in source header"}
    label = _slug(" ".join(filter(None, [str(table.get("sheet") or ""), " ".join(table["headers"])])))
    values = [parse_number(row[fields[measure_key]] if fields[measure_key] < len(row) else None) for row in table["rows"]]
    values = [value for value in values if value is not None]
    timestamps = [parse_date(row[fields["date"]] if fields["date"] < len(row) else None) for row in table["rows"]]
    timestamps = sorted(item for item in timestamps if item is not None)
    regular_subdaily = False
    if len(timestamps) >= 3:
        deltas = [(right - left).total_seconds() / 60 for left, right in zip(timestamps, timestamps[1:])]
        typical = median(deltas)
        regular_subdaily = 5 <= typical <= 240 and max(abs(value - typical) for value in deltas) <= max(1, typical * .05)
    if measure_key == "power":
        return {"status": "MATERIAL_AMBIGUITY", "unit": None, "measurement_kind": "power", "basis": "power-like column lacks a physical unit"}
    if "kwh" in label or ("conso" in header and regular_subdaily and role == "ENERGY_INTERVAL_SERIES"):
        return {"status": "AUTO_FIX_SAFE", "unit": "kwh", "measurement_kind": "energy_per_interval", "basis": "energy context plus regular sub-daily series"}
    if values and all(value >= 0 for value in values):
        return {"status": "MATERIAL_AMBIGUITY", "unit": None, "measurement_kind": "unknown", "basis": "unlabeled non-negative measurement supports competing kW/kWh interpretations"}
    return {"status": "UNUSABLE", "unit": None, "measurement_kind": "unknown", "basis": "no usable numeric measurement values"}


def _normalise_table(table: dict[str, Any], *, dataset_id: str, role: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    """Parse safely; never repairs suspicious observations or conflicting duplicates."""

    fields = table["field_indices"]
    measurement = "power" if "power" in fields else "energy"
    unit = _inferred_unit(table, role=role)
    transformations: list[dict[str, Any]] = []
    records: list[dict[str, Any]] = []
    invalid_timestamps = 0
    total_rows = 0
    decimal_comma_values = 0
    for offset, row in enumerate(table["rows"], start=table["header_row"] + 1):
        text_row = " ".join(str(value or "") for value in row).strip()
        if any(token in {"total", "totaux", "sous", "subtotal"} for token in _slug(text_row).split("_")):
            total_rows += 1
            continue
        timestamp = parse_date(row[fields["date"]] if fields["date"] < len(row) else None)
        raw_value = row[fields[measurement]] if fields[measurement] < len(row) else None
        value = parse_number(raw_value)
        if isinstance(raw_value, str) and "," in raw_value:
            decimal_comma_values += 1
        if timestamp is None:
            invalid_timestamps += 1
            continue
        record = {"timestamp": timestamp.isoformat(), "source_row": offset, "measurement": value}
        for semantic in ("production", "temperature", "tariff", "shift", "product_type", "production_active"):
            index = fields.get(semantic)
            if index is None or index >= len(row):
                continue
            raw = row[index]
            record[semantic] = parse_number(raw) if semantic in {"production", "temperature", "tariff"} else raw
        records.append(record)
    if total_rows:
        transformations.append({"status": "AUTO_FIX_SAFE", "operation": "exclude_total_rows", "count": total_rows, "reason": "explicit total label, not an observation"})
    if decimal_comma_values:
        transformations.append({"status": "AUTO_FIX_SAFE", "operation": "normalize_decimal_comma", "count": decimal_comma_values, "reason": "standard numeric parsing of French decimal notation"})
    if invalid_timestamps:
        transformations.append({"status": "FLAG_ONLY", "operation": "invalid_timestamp_rows", "count": invalid_timestamps, "reason": "rows preserved in raw source and excluded only from timestamped table"})
    if unit["status"] != "AUTO_FIX_SAFE":
        transformations.append({"status": unit["status"], "operation": "measurement_unit", "reason": unit["basis"]})
    elif unit["basis"] != "explicit unit in source header":
        transformations.append({"status": "AUTO_FIX_SAFE", "operation": "infer_unit", "value": unit["unit"], "reason": unit["basis"]})
    by_timestamp: dict[str, list[dict[str, Any]]] = {}
    for record in records:
        by_timestamp.setdefault(record["timestamp"], []).append(record)
    strict_duplicates = 0
    conflicts = 0
    retained: list[dict[str, Any]] = []
    for timestamp, group in by_timestamp.items():
        reference = group[0]
        comparable = [{key: value for key, value in item.items() if key != "source_row"} for item in group]
        if len(group) > 1 and all(item == comparable[0] for item in comparable[1:]):
            retained.append(reference)
            strict_duplicates += len(group) - 1
        elif len(group) > 1:
            retained.extend(group)
            conflicts += len(group) - 1
        else:
            retained.append(reference)
    if strict_duplicates:
        transformations.append({"status": "AUTO_FIX_SAFE", "operation": "remove_strict_duplicates", "count": strict_duplicates, "reason": "all measured fields identical"})
    if conflicts:
        transformations.append({"status": "MATERIAL_AMBIGUITY", "operation": "conflicting_duplicates", "count": conflicts, "reason": "same timestamp has different measurements; no value selected"})
    ordered = sorted(retained, key=lambda item: item["timestamp"])
    timestamps = [datetime.fromisoformat(item["timestamp"]) for item in ordered]
    deltas = [(right - left).total_seconds() / 60 for left, right in zip(timestamps, timestamps[1:])]
    frequency = median(deltas) if deltas else None
    missing_intervals = 0
    frequency_change = False
    coverage_ratio = None
    if frequency and frequency > 0:
        missing_intervals = sum(max(0, round(delta / frequency) - 1) for delta in deltas if delta > frequency * 1.5)
        frequency_change = any(abs(delta - frequency) > max(1, frequency * .05) and delta <= frequency * 1.5 for delta in deltas)
        if timestamps and not frequency_change:
            expected = round((timestamps[-1] - timestamps[0]).total_seconds() / 60 / frequency) + 1
            coverage_ratio = min(1.0, len(ordered) / expected) if expected else None
    suspicious = [item["source_row"] for item in ordered if item.get("production") is not None and item["production"] < 0]
    measurement_negative = [item["source_row"] for item in ordered if item.get("measurement") is not None and item["measurement"] < 0]
    if suspicious:
        transformations.append({"status": "FLAG_ONLY", "operation": "negative_production", "source_rows": suspicious, "reason": "not rewritten as a presumed entry error"})
    if measurement_negative:
        transformations.append({"status": "FLAG_ONLY", "operation": "negative_measurement", "source_rows": measurement_negative, "reason": "not silently removed"})
    quality = {
        "usable_time_span": {"start": ordered[0]["timestamp"] if ordered else None, "end": ordered[-1]["timestamp"] if ordered else None},
        "rows_parsed": len(records), "rows_retained": len(ordered), "sampling_frequency_minutes": frequency,
        "coverage_ratio": coverage_ratio, "missing_intervals": missing_intervals, "frequency_change_detected": frequency_change,
        "strict_duplicates_removed": strict_duplicates, "conflicting_duplicates": conflicts,
        "suspicious_values": {"negative_production_rows": suspicious, "negative_measurement_rows": measurement_negative},
        "production_coverage_ratio": (sum(item.get("production") is not None for item in ordered) / len(ordered)) if ordered else None,
        "unit_interpretation": unit, "timezone_confidence": "not_declared", "known_limitations": [item["reason"] for item in transformations if item["status"] in {"MATERIAL_AMBIGUITY", "UNUSABLE"}],
    }
    return ordered, transformations, quality



def _normalise_context_table(table: dict[str, Any], *, dataset_id: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    """Normalise une table contexte (production, météo) sans la faire passer pour une mesure énergie."""

    fields = table["field_indices"]
    records: list[dict[str, Any]] = []
    invalid = 0
    for offset, row in enumerate(table["rows"], start=table["header_row"] + 1):
        timestamp = parse_date(row[fields["date"]] if fields["date"] < len(row) else None)
        if timestamp is None:
            invalid += 1
            continue
        item = {"timestamp": timestamp.isoformat(), "source_row": offset}
        for semantic, index in fields.items():
            if semantic == "date" or index >= len(row):
                continue
            raw = row[index]
            item[semantic] = parse_number(raw) if semantic in {"production", "temperature", "tariff"} else raw
        records.append(item)
    ordered = sorted(records, key=lambda item: item["timestamp"])
    transformations = ([{"status": "FLAG_ONLY", "operation": "invalid_timestamp_rows", "count": invalid, "reason": "context rows with invalid dates remain available in raw source"}] if invalid else [])
    quality = {
        "usable_time_span": {"start": ordered[0]["timestamp"] if ordered else None, "end": ordered[-1]["timestamp"] if ordered else None},
        "rows_parsed": len(records), "rows_retained": len(ordered), "sampling_frequency_minutes": None,
        "missing_intervals": None, "frequency_change_detected": False, "strict_duplicates_removed": 0,
        "conflicting_duplicates": 0, "suspicious_values": {}, "unit_interpretation": {
            "status": "FLAG_ONLY", "unit": None, "measurement_kind": "context_only",
            "basis": "context table is not an energy/power measurement",
        }, "timezone_confidence": "not_declared", "known_limitations": [],
    }
    return ordered, transformations, quality


def _write_context_json(path: Path, records: list[dict[str, Any]]) -> None:
    _write_json(path, {"schema_version": 1, "records": records})

def _write_normalized_csv(path: Path, records: list[dict[str, Any]], *, measurement_kind: str, unit: str) -> None:
    fields = ["timestamp", f"{measurement_kind}_{unit}", "production", "temperature", "tariff", "shift", "product_type", "production_active", "source_row"]
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for record in records:
            writer.writerow({
                "timestamp": record.get("timestamp"), f"{measurement_kind}_{unit}": record.get("measurement"),
                "production": record.get("production"), "temperature": record.get("temperature"),
                "tariff": record.get("tariff"), "shift": record.get("shift"),
                "product_type": record.get("product_type"), "production_active": record.get("production_active"),
                "source_row": record.get("source_row"),
            })


def _inspect_file(path: Path, *, artifact_id: str, raw_relative: str) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    suffix = path.suffix.lower()
    tables: list[dict[str, Any]] = []
    parser_status = "not_parsed"
    parse_error: str | None = None
    try:
        if suffix == ".csv":
            candidate = _candidate_table(_read_csv_rows(path), sheet=None)
            if candidate:
                tables.append(candidate)
            parser_status = "parsed"
        elif suffix == ".xlsx":
            for sheet, rows in _read_xlsx_sheets(path).items():
                candidate = _candidate_table(rows, sheet=sheet)
                if candidate:
                    tables.append(candidate)
            parser_status = "parsed"
        elif suffix == ".xls":
            parser_status = "unreadable"
            parse_error = "XLS legacy non lu sans dépendance optionnelle; fichier conservé pour conversion ou lecture ultérieure."
        elif suffix in _TEXT_EXTENSIONS:
            parser_status = "metadata_only"
        else:
            parser_status = "unsupported"
    except (OSError, ValueError, UnicodeError) as exc:
        parser_status = "unreadable"
        parse_error = str(exc)
    if tables:
        role, evidence = _role_from_table(tables[0])
        usability = "usable"
    else:
        role, evidence = _role_from_name(path)
        usability = "unreadable" if parser_status == "unreadable" else "irrelevant" if role == "IRRELEVANT" else "partially_usable"
    inventory = {
        "artifact_id": artifact_id, "original_filename": path.name, "raw_relative_path": raw_relative,
        "file_type": suffix.lstrip(".") or "unknown", "mime_type": mimetypes.guess_type(path.name)[0],
        "size_bytes": path.stat().st_size, "sha256": _sha256(path), "ingested_at_utc": _utc_now(),
        "parsing_status": parser_status, "probable_role": role, "role_evidence": evidence,
        "usability": usability, "parse_error": parse_error, "extracted_dataset_ids": [],
    }
    return inventory, tables


def ingest_client_drop(source_directory: str | Path, case_directory: str | Path) -> dict[str, Any]:
    """Copie un drop client sans modifier les originaux puis produit son paquet canonique.

    La détection est volontairement descriptive : les rôles et les unités sont des
    observations/inférences tracées, pas des conclusions énergétiques.
    """

    source = Path(source_directory).resolve()
    case = Path(case_directory).resolve()
    if not source.is_dir():
        raise ValueError(f"Dossier client introuvable: {source}")
    if not (case / "case_manifest.json").is_file():
        raise ValueError("Le dossier cible doit avoir été créé par create_client_case.")
    inventory_path = case / "evidence" / "intake_inventory.json"
    if inventory_path.exists():
        raise FileExistsError("Un intake existe déjà; il ne sera pas réécrit silencieusement.")
    inventory: list[dict[str, Any]] = []
    datasets: list[dict[str, Any]] = []
    seen_hashes: dict[str, str] = {}
    for index, original in enumerate(_safe_relative_files(source), start=1):
        relative = original.relative_to(source)
        artifact_id = f"ART-{index:03d}-{_sha256(original)[:12]}"
        raw_target = case / "raw" / f"{artifact_id}_{relative.name}"
        shutil.copy2(original, raw_target)
        record, tables = _inspect_file(raw_target, artifact_id=artifact_id, raw_relative=str(relative))
        if record["sha256"] in seen_hashes:
            record["duplicate_of_artifact_id"] = seen_hashes[record["sha256"]]
        else:
            seen_hashes[record["sha256"]] = artifact_id
        for table_index, table in enumerate(tables, start=1):
            role, role_evidence = _role_from_table(table)
            dataset_id = f"DS-{artifact_id.split('-', 2)[1]}-{table_index:02d}"
            is_measurement = role in {"ENERGY_INTERVAL_SERIES", "ENERGY_MONTHLY", "POWER_SERIES"}
            if is_measurement:
                records, transformations, quality = _normalise_table(table, dataset_id=dataset_id, role=role)
            else:
                records, transformations, quality = _normalise_context_table(table, dataset_id=dataset_id)
            unit = quality["unit_interpretation"]
            analytical_usable = is_measurement and unit["status"] == "AUTO_FIX_SAFE" and quality["conflicting_duplicates"] == 0 and bool(records)
            normalized_file = None
            if analytical_usable:
                kind = "power" if unit["measurement_kind"] == "power" else "energy"
                normalized_file = f"normalized/{dataset_id}.csv"
                _write_normalized_csv(case / normalized_file, records, measurement_kind=kind, unit=unit["unit"])
            elif records:
                normalized_file = f"normalized/{dataset_id}.json"
                _write_context_json(case / normalized_file, records)
            measurement_index = table["field_indices"].get("power", table["field_indices"].get("energy"))
            field_lineage = {"timestamp": {"source_artifact_id": artifact_id, "sheet": table["sheet"], "source_column": table["headers"][table["field_indices"]["date"]], "normalization": "parse_date"}}
            if measurement_index is not None:
                field_lineage["measurement"] = {"source_artifact_id": artifact_id, "sheet": table["sheet"], "source_column": table["headers"][measurement_index], "normalization": "parse_number", "unit_interpretation": unit}
            if "production" in table["field_indices"]:
                field_lineage["production"] = {"source_artifact_id": artifact_id, "sheet": table["sheet"], "source_column": table["headers"][table["field_indices"]["production"]], "normalization": "parse_number"}
            dataset = {
                "dataset_id": dataset_id, "source_artifact_id": artifact_id, "role": role,
                "role_evidence": role_evidence, "sheet": table["sheet"], "header_row": table["header_row"],
                "source_headers": table["headers"], "source_columns": {name: table["headers"][column] for name, column in table["field_indices"].items()},
                "normalized_file": normalized_file, "analytically_usable": analytical_usable,
                "transformations": transformations, "data_quality": quality, "field_lineage": field_lineage,
            }
            datasets.append(dataset)
            record["extracted_dataset_ids"].append(dataset_id)
        inventory.append(record)
    if not inventory:
        raise ValueError("Le dossier client ne contient aucun fichier exploitable à inventorier.")
    _write_json(inventory_path, {"schema_version": 1, "case_id": case.name, "created_at_utc": _utc_now(), "artifacts": inventory})
    _write_json(case / "evidence" / "dataset_provenance.json", {"schema_version": 1, "datasets": datasets})
    canonical = _canonical_case(case, inventory, datasets)
    _write_json(case / "derived" / "canonical_case.json", canonical)
    _write_json(case / "derived" / "data_quality_summary.json", _quality_summary(datasets, inventory))
    _write_json(case / "investigation" / "question_batch.json", _empty_question_batch(case.name))
    state = _read_json(case / "investigation" / "case_state.json")
    state.update({"status": "ready_for_codex_first_pass", "available_datasets": [item["dataset_id"] for item in datasets], "material_ambiguities": canonical["unresolved_material_ambiguities"], "history": state["history"] + [{"at_utc": _utc_now(), "action": "intake_completed", "inventory": "evidence/intake_inventory.json"}]})
    _write_json(case / "investigation" / "case_state.json", state)
    _write_first_pass_brief(case, canonical)
    manifest = _read_json(case / "case_manifest.json")
    manifest["status"] = "ready_for_codex_first_pass"
    manifest["intake_completed_at_utc"] = _utc_now()
    _write_json(case / "case_manifest.json", manifest)
    engine_handoff(case)
    return canonical


def _read_json(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"JSON invalide: {path}")
    return payload


def _quality_summary(datasets: list[dict[str, Any]], inventory: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        "schema_version": 1,
        "datasets": [{
            "dataset_id": item["dataset_id"], "role": item["role"], "usable_time_span": item["data_quality"]["usable_time_span"],
            "sampling_frequency_minutes": item["data_quality"]["sampling_frequency_minutes"], "coverage_ratio": item["data_quality"].get("coverage_ratio"),
            "missing_intervals": item["data_quality"]["missing_intervals"], "production_coverage_ratio": item["data_quality"].get("production_coverage_ratio"),
            "timezone_confidence": item["data_quality"].get("timezone_confidence"), "conflicting_duplicates": item["data_quality"]["conflicting_duplicates"], "unit_interpretation": item["data_quality"]["unit_interpretation"],
            "known_limitations": item["data_quality"]["known_limitations"],
        } for item in datasets],
        "supporting_document_coverage": Counter(item["probable_role"] for item in inventory),
        "global_score": None,
        "note": "Aucun score global: la qualité limite seulement les conclusions qu'elle affecte.",
    }


def _canonical_case(case: Path, inventory: list[dict[str, Any]], datasets: list[dict[str, Any]]) -> dict[str, Any]:
    ambiguities = []
    for dataset in datasets:
        for item in dataset["transformations"]:
            if item["status"] == "MATERIAL_AMBIGUITY":
                ambiguities.append({"dataset_id": dataset["dataset_id"], "issue": item["operation"], "reason": item["reason"], "blocking_for": "quantitative_energy_interpretation"})
    by_role: dict[str, list[str]] = {}
    for dataset in datasets:
        by_role.setdefault(dataset["role"], []).append(dataset["dataset_id"])
    for item in inventory:
        if item["probable_role"] not in by_role:
            by_role.setdefault(item["probable_role"], []).append(item["artifact_id"])
    return {
        "schema_version": 1, "case_id": case.name, "created_at_utc": _utc_now(),
        "purpose": "Paquet descriptif pour investigation Codex; aucune anomalie, cause, action ou question n'est décidée ici.",
        "available_datasets": [{key: value for key, value in item.items() if key in {"dataset_id", "role", "normalized_file", "analytically_usable", "data_quality", "field_lineage"}} for item in datasets],
        "evidence_by_role": by_role,
        "known_facts": [{"kind": "source_observation", "value": f"{item['original_filename']} classé {item['probable_role']}", "artifact_id": item["artifact_id"]} for item in inventory],
        "inferred_facts": [{"dataset_id": item["dataset_id"], "measurement": item["data_quality"]["unit_interpretation"]} for item in datasets if item["data_quality"]["unit_interpretation"]["status"] == "AUTO_FIX_SAFE"],
        "unresolved_material_ambiguities": ambiguities,
        "agent_decisions_required": ["anomaly relevance", "hypotheses", "alternative explanations", "next calculation", "client question necessity", "finding status", "intervention"],
        "no_finding_is_valid": True,
        "ground_truth_available": False,
    }


def _write_first_pass_brief(case: Path, canonical: dict[str, Any]) -> None:
    datasets = canonical["available_datasets"]
    (case / "investigation" / "CODEX_FIRST_PASS_BRIEF.md").write_text("\n".join([
        "# Premier passage Codex — dossier client", "",
        "Ce paquet est descriptif. Les classifications de fichiers et transformations ne sont pas un diagnostic.", "",
        "- Commencer par `derived/canonical_case.json`, `derived/data_quality_summary.json` et `evidence/dataset_provenance.json`.",
        f"- Jeux disponibles: {', '.join(item['dataset_id'] for item in datasets) or 'aucun'}.",
        f"- Ambiguïtés matérielles ouvertes: {len(canonical['unresolved_material_ambiguities'])}.",
        "- Exploiter les données existantes avant toute question client.",
        "- Python mesure; Codex choisit les pistes, les hypothèses, les contrôles et les conclusions.",
        "- Une absence de finding est un résultat valide. Ne pas inventer une opportunité.",
        "- Ne demander une information que si elle modifie matériellement une décision; par défaut, zéro question.",
        "- Ne pas écrire de chaîne de pensée détaillée: consigner des résumés auditables, preuves et contre-explications.",
    ]) + "\n", encoding="utf-8")


def _empty_question_batch(case_id: str) -> dict[str, Any]:
    return {"schema_version": 1, "case_id": case_id, "batch_status": "not_needed_yet", "requests": [], "client_facing_text": "Aucune question client nécessaire à ce stade."}


def publish_question_batch(case_directory: str | Path, requests: list[dict[str, Any]], *, exceptional_second_batch_reason: str | None = None) -> dict[str, Any]:
    """Publie une demande fournie par Codex, sans en générer une automatiquement."""

    case = Path(case_directory)
    if not 1 <= len(requests) <= MAX_DEFAULT_QUESTIONS:
        raise ValueError("Un batch client comporte normalement de 1 à 3 demandes ciblées.")
    path = case / "investigation" / "question_batch.json"
    current = _read_json(path)
    if current.get("requests") and not exceptional_second_batch_reason:
        raise ValueError("Un second batch exige une justification matérielle explicite.")
    if current.get("requests"):
        archive_directory = case / "investigation" / "question_batch_history"
        archive_directory.mkdir(exist_ok=True)
        index = 1
        while (archive_directory / f"batch_{index:03d}.json").exists():
            index += 1
        _write_json(archive_directory / f"batch_{index:03d}.json", current)
    normalized: list[dict[str, Any]] = []
    seen: set[str] = set()
    for position, request in enumerate(requests, start=1):
        request_id = str(request.get("request_id") or f"REQ-{position:02d}")
        if request_id in seen:
            raise ValueError("request_id dupliqué.")
        seen.add(request_id)
        request_type = request.get("request_type")
        importance = request.get("importance", "NON_BLOCKING")
        for field in ("client_question", "internal_reason", "target_role", "decision_impact", "expected_effort"):
            if not str(request.get(field, "")).strip():
                raise ValueError(f"{request_id}: {field} est requis.")
        if request_type not in REQUEST_TYPES or request_type == "INFER_AUTOMATICALLY":
            raise ValueError(f"{request_id}: type de demande client invalide.")
        if importance not in REQUEST_IMPORTANCE:
            raise ValueError(f"{request_id}: importance invalide.")
        text = str(request["client_question"]).strip()
        if _slug(text) in {"plus_de_donnees", "davantage_de_donnees"} or len(text) < 12:
            raise ValueError(f"{request_id}: demande client trop vague.")
        normalized.append({
            "request_id": request_id, "request_type": request_type, "client_question": text,
            "internal_reason": str(request["internal_reason"]).strip(), "target_role": str(request["target_role"]).strip(),
            "hypotheses_distinguished": list(request.get("hypotheses_distinguished") or []), "decision_impact": str(request["decision_impact"]).strip(),
            "expected_effort": str(request["expected_effort"]).strip(), "importance": importance,
            "analysis_can_continue_without_answer": importance != "BLOCKING", "created_at_utc": _utc_now(),
        })
    payload = {"schema_version": 1, "case_id": case.name, "batch_status": "published", "requests": normalized, "client_facing_text": "\n".join(f"- {item['client_question']}" for item in normalized), "exceptional_second_batch_reason": exceptional_second_batch_reason}
    _write_json(path, payload)
    state_path = case / "investigation" / "case_state.json"
    state = _read_json(state_path)
    state["unresolved_requests"] = normalized
    state["history"].append({"at_utc": _utc_now(), "action": "question_batch_published", "request_ids": [item["request_id"] for item in normalized]})
    _write_json(state_path, state)
    return payload


def record_structured_findings(case_directory: str | Path, findings: list[dict[str, Any]], *, no_finding: dict[str, Any] | None = None) -> dict[str, Any]:
    """Conserve les conclusions structurées choisies par Codex; ne les déduit jamais."""

    case = Path(case_directory)
    if findings and no_finding:
        raise ValueError("Un cas ne peut pas publier simultanément findings matériels et no_finding.")
    if not findings and not no_finding:
        raise ValueError("Fournissez des findings ou une conclusion no_finding structurée.")
    for finding in findings:
        for field in ("finding_id", "observation", "status", "possible_explanations", "provenance", "recommended_next_analytical_step"):
            if not finding.get(field):
                raise ValueError(f"Finding incomplet: {field} requis.")
    if no_finding:
        for field in ("what_was_analyzed", "usable_period", "operating_regimes", "limitations", "monitoring_baseline_meaningful"):
            if field not in no_finding:
                raise ValueError(f"no_finding incomplet: {field} requis.")
    payload = {"schema_version": 1, "created_at_utc": _utc_now(), "findings": findings, "no_finding": no_finding, "decision_source": "Codex investigation; deterministic pipeline did not decide findings."}
    _write_json(case / "investigation" / "structured_findings.json", payload)
    state_path = case / "investigation" / "case_state.json"
    state = _read_json(state_path)
    state["findings"] = findings
    state["current_decision_status"] = "no_material_finding" if no_finding else "findings_recorded"
    state["history"].append({"at_utc": _utc_now(), "action": "structured_findings_recorded", "count": len(findings), "no_finding": bool(no_finding)})
    _write_json(state_path, state)
    return payload


def record_client_answers(case_directory: str | Path, answers: list[dict[str, Any]]) -> dict[str, Any]:
    """Enregistre des réponses reçues sans autoriser leur réécriture."""

    case = Path(case_directory)
    batch_path = case / "investigation" / "question_batch.json"
    batch = _read_json(batch_path)
    requests = {item["request_id"]: item for item in batch.get("requests", [])}
    if not answers:
        raise ValueError("Au moins une réponse est requise.")
    state_path = case / "investigation" / "case_state.json"
    state = _read_json(state_path)
    answered_ids = {item.get("request_id") for item in state.get("client_answers", [])}
    recorded: list[dict[str, Any]] = []
    for answer in answers:
        request_id = answer.get("request_id")
        if request_id not in requests:
            raise ValueError(f"Réponse pour une demande inconnue: {request_id}.")
        if request_id in answered_ids:
            raise ValueError(f"La réponse {request_id} est déjà enregistrée et ne sera pas réécrite.")
        for field in ("answer", "provided_by_role", "source_or_evidence"):
            if not str(answer.get(field, "")).strip():
                raise ValueError(f"{request_id}: {field} est requis.")
        item = {
            "request_id": request_id,
            "answer": str(answer["answer"]).strip(),
            "provided_by_role": str(answer["provided_by_role"]).strip(),
            "source_or_evidence": str(answer["source_or_evidence"]).strip(),
            "received_at_utc": _utc_now(),
        }
        requests[request_id]["status"] = "answered"
        requests[request_id]["answer"] = item
        recorded.append(item)
    state["client_answers"].extend(recorded)
    state["unresolved_requests"] = [item for item in batch["requests"] if item.get("status") != "answered"]
    state["history"].append({"at_utc": _utc_now(), "action": "client_answers_recorded", "request_ids": [item["request_id"] for item in recorded]})
    _write_json(batch_path, batch)
    _write_json(state_path, state)
    return {"schema_version": 1, "recorded_answers": recorded, "analysis_may_continue": any(item.get("importance") != "BLOCKING" for item in state["unresolved_requests"])}


def engine_handoff(case_directory: str | Path) -> dict[str, Any]:
    """Retourne les jeux normalisés que Codex peut choisir d'investiguer.

    Aucune sélection ou analyse n'est faite ici : avec plusieurs compteurs, le
    périmètre pertinent est une décision de l'analyste.
    """

    case = Path(case_directory)
    canonical = _read_json(case / "derived" / "canonical_case.json")
    candidates = []
    for dataset in canonical["available_datasets"]:
        if dataset.get("analytically_usable") and dataset.get("normalized_file"):
            candidates.append({
                "dataset_id": dataset["dataset_id"],
                "role": dataset["role"],
                "normalized_input": dataset["normalized_file"],
                "provenance": "evidence/dataset_provenance.json",
                "suggested_existing_command": f"python investigate.py {dataset['normalized_file']} --output-dir investigation/{dataset['dataset_id']}",
            })
    payload = {
        "schema_version": 1,
        "purpose": "Choix explicite de Codex avant les calculs; aucun jeu n'est privilégié automatiquement.",
        "analytically_usable_inputs": candidates,
        "not_ready_without_resolution": canonical["unresolved_material_ambiguities"],
        "agent_decision_required": "Choisir le périmètre, les tests et l'ordre d'investigation.",
    }
    _write_json(case / "investigation" / "engine_handoff.json", payload)
    return payload

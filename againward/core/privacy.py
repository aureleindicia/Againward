"""Privacy gate contracts for real INDICIA client workspaces.

Codex is the first semantic reader of ``incoming/``.  This module deliberately
does not classify an incoming drop before that review.  It validates Codex's
structured decision afterwards, checks obvious residual PII/secrets, verifies
that industrial series were not changed, promotes an approved source, and
records an audit manifest that never contains removed values.
"""
from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import shutil
import tempfile
import unicodedata
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

from .workflow_paths import resolve_case_layout
from .privacy_rules import current_preservation_policy, with_preservation_policy
from .privacy_inspection import INSPECTION_VERSION, inspect_document


POLICY_VERSION = "indicia-privacy-policy-v1.1"
REVIEW_SCHEMA = "indicia-codex-privacy-review-v1"
MANIFEST_SCHEMA = "indicia-privacy-manifest-v1"
RETENTION_SCHEMA = "indicia-retention-policy-v1"
PURGE_SCHEMA = "indicia-purge-receipt-v1"
PRIVACY_STATES = {
    "AWAITING_PRIVACY_REVIEW", "PRIVACY_CLEARED", "PRIVACY_BLOCKED", "PURGED",
}

_TABULAR_EXTENSIONS = {".csv", ".xlsx"}
_TEXT_EXTENSIONS = {".txt", ".md", ".json", ".log", ".yaml", ".yml"}
_EMAIL = re.compile(r"(?<![\w.+-])[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}(?![\w-])", re.I)
_PERSONAL_EMAIL_LABEL = re.compile(r"(?i)\b(?:personal|private|personnel|priv[eé])\s+(?:e-?mail|courriel)\b")
_PHONE = re.compile(r"(?<!\d)(?:\+33|0033|0)[ .()-]?[1-9](?:[ .()-]?\d{2}){4}(?!\d)")
_SECRET_PATTERNS = (
    re.compile(r"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----"),
    re.compile(r"\bAKIA[0-9A-Z]{16}\b"),
    re.compile(r"\b(?:sk|rk|pk)-(?:live|test|proj)?-?[A-Za-z0-9_-]{16,}\b"),
    re.compile(r"(?i)\b(?:api[_ -]?key|access[_ -]?token|client[_ -]?secret|password|passwd|mot[_ -]?de[_ -]?passe)\b\s*[:=]\s*[^\s,;]{6,}"),
    re.compile(r"(?i)https?://[^\s]+[?&](?:token|api_key|access_token|key)=[^&\s]+"),
)
_MEDICAL = re.compile(r"(?i)\b(?:arr[eê]t maladie|diagnostic m[eé]dical|medical diagnosis|medical record|handicap|pathologie|traitement m[eé]dical|dossier m[eé]dical)\b")
_HR_SENSITIVE = re.compile(r"(?i)\b(?:sanction disciplinaire|avertissement rh|licenciement|entretien disciplinaire|plainte harc[eè]lement|[eé]valuation individuelle|disciplinary action|personnel file)\b")
_IDENTITY_DOCUMENT = re.compile(r"(?i)\b(?:passport(?: number| scan| copy)?|passeport|carte (?:nationale )?d.identit[eé]|identity card|driver.s license copy)\b")
_RESIDENTIAL_ADDRESS = re.compile(r"(?i)\b(?:home address|residential address|adresse personnelle|adresse du domicile)\b")
_HARD_PRIVACY_CATEGORIES = frozenset({"MEDICAL_DATA", "HR_SENSITIVE", "IDENTITY_DOCUMENT", "UNRELATED_PERSONAL_RECORD"})
_PSEUDONYM = re.compile(r"^(?:OPERATOR|TECHNICIAN|EMPLOYEE|PERSON|WORKER|STAFF|ID)_[0-9A-F]{3,32}$", re.I)
_SAFE_CODE = re.compile(r"^[A-Z][A-Z0-9_]{1,63}$")
_FILE_ID = re.compile(r"^FILE-[0-9]{3,9}$")
_MANAGED_CASE_PARENT_NAMES = {"workspaces", "client_cases"}

_PERSONAL_HEADER_TOKENS = {
    "first_name", "firstname", "prenom", "last_name", "lastname", "nom_personne",
    "full_name", "fullname", "employee_name", "operator_name", "technician_name",
    "technician", "technicien", "operator", "operateur", "employee", "salarie",
    "matricule", "employee_id", "operator_id", "badge_id", "user_id", "email",
    "e_mail", "mail", "phone", "telephone", "mobile", "address", "adresse",
    "personal_address", "rh", "medical", "medical_info", "api_key", "access_token",
    "client_secret", "password", "passwd", "credential", "credentials",
}
_DIRECT_REMOVE_HEADER_TOKENS = {
    "first_name", "firstname", "prenom", "last_name", "lastname", "nom_personne",
    "full_name", "fullname", "employee_name", "operator_name", "technician_name",
    "email", "e_mail", "mail", "phone", "telephone", "mobile", "address",
    "adresse", "personal_address", "medical", "medical_info", "rh", "api_key",
    "access_token", "client_secret", "password", "passwd", "credential", "credentials",
}
_RELATION_HEADER_TOKENS = {
    "technician", "technicien", "operator", "operateur", "employee", "salarie",
    "matricule", "employee_id", "operator_id", "badge_id", "user_id",
}
_FREE_TEXT_HEADER_TOKENS = {"comment", "comments", "commentaire", "commentaires", "note", "notes", "free_text", "texte_libre"}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _slug(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "_", text.casefold()).strip("_")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _read_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ValueError(f"Fichier requis absent: {path}.") from exc
    except json.JSONDecodeError as exc:
        raise ValueError(f"JSON invalide dans {path}: {exc}.") from exc
    if not isinstance(value, dict):
        raise ValueError(f"Objet JSON attendu: {path}.")
    return value


def _atomic_json(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    encoded = json.dumps(value, ensure_ascii=False, indent=2) + "\n"
    temporary: str | None = None
    try:
        with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=path.parent, prefix=f".{path.name}.", delete=False) as handle:
            temporary = handle.name
            handle.write(encoded)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        if temporary:
            Path(temporary).unlink(missing_ok=True)


def _safe_files(root: Path) -> list[Path]:
    if root.is_symlink():
        raise ValueError(f"Lien symbolique interdit dans les données client: {root}.")
    if not root.is_dir():
        return []
    result: list[Path] = []
    for path in sorted(root.rglob("*")):
        if path.is_symlink():
            raise ValueError(f"Lien symbolique interdit dans les données client: {path}.")
        if path.is_file():
            result.append(path)
    return result


def _within(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
        return True
    except ValueError:
        return False


def _conventional_case_root(path: str | Path) -> Path | None:
    """Recognize an unregistered case below an official managed-case root."""

    current = Path(path).expanduser().resolve(strict=False)
    if current.is_file() or current.suffix:
        current = current.parent
    for candidate in (current, *current.parents):
        if candidate.parent.name in _MANAGED_CASE_PARENT_NAMES:
            return candidate
    return None


def privacy_requirement(case_directory: str | Path) -> dict[str, Any]:
    """Describe whether a recognized case requires privacy clearance."""

    layout = resolve_case_layout(case_directory)
    root: Path = layout["case_root"]
    manifest_path = root / ("workspace.json" if layout["layout"] == "STANDARD_WORKSPACE" else "case_manifest.json")
    if not manifest_path.is_file():
        managed_root = _conventional_case_root(case_directory)
        if layout["layout"] == "GOAL_A_CASE" or managed_root is not None:
            return {
                "required": True,
                "legacy": True,
                "layout": layout["layout"],
                "case_root": managed_root or root,
            }
        return {"required": False, "legacy": False, "layout": layout["layout"], "case_root": root}
    manifest = _read_json(manifest_path)
    privacy = manifest.get("privacy")
    if not isinstance(privacy, dict) or "required" not in privacy:
        return {"required": True, "legacy": True, "layout": layout["layout"], "case_root": root}
    case_kind = manifest.get("case_kind")
    if case_kind not in {"REAL_CLIENT", "SYNTHETIC"}:
        return {"required": True, "legacy": True, "layout": layout["layout"], "case_root": root}
    if case_kind == "REAL_CLIENT" and privacy.get("required") is not True:
        return {"required": True, "legacy": True, "layout": layout["layout"], "case_root": root}
    if case_kind == "REAL_CLIENT" and privacy.get("policy_version") != POLICY_VERSION:
        return {"required": True, "legacy": True, "layout": layout["layout"], "case_root": root}
    if case_kind == "SYNTHETIC" and privacy.get("required") is not False:
        return {"required": True, "legacy": True, "layout": layout["layout"], "case_root": root}
    return {
        "required": bool(privacy.get("required")),
        "legacy": False,
        "synthetic": case_kind == "SYNTHETIC",
        "layout": layout["layout"],
        "case_root": root,
    }


def privacy_manifest_path(case_directory: str | Path) -> Path:
    return Path(privacy_requirement(case_directory)["case_root"]) / "privacy" / "privacy_manifest.json"


def inspect_privacy_status(case_directory: str | Path) -> dict[str, Any]:
    requirement = privacy_requirement(case_directory)
    if not requirement["required"]:
        return {"state": "NOT_REQUIRED", "approved_for_analysis": True, **requirement}
    if requirement["legacy"]:
        return {"state": "PRIVACY_MIGRATION_REQUIRED", "approved_for_analysis": False, **requirement}
    path = privacy_manifest_path(case_directory)
    if not path.is_file():
        return {"state": "AWAITING_PRIVACY_REVIEW", "approved_for_analysis": False, **requirement}
    manifest = _read_json(path)
    if manifest.get("policy_version") != POLICY_VERSION or manifest.get("schema_version") != MANIFEST_SCHEMA:
        return {
            "state": "PRIVACY_MIGRATION_REQUIRED",
            "approved_for_analysis": False,
            "manifest": str(path),
            **requirement,
        }
    case_meta = requirement["case_root"] / ("workspace.json" if requirement["layout"] == "STANDARD_WORKSPACE" else "case_manifest.json")
    if manifest.get("status") in {"PASS", "SANITIZED"} and _read_json(case_meta).get("domain") == "rental":
        from againward.entrypoints import get_case_domain
        current_risk = get_case_domain(requirement["case_root"]).privacy_preservation.risk.policy_id
        if manifest.get("risk_policy_version") != current_risk:
            return {"state": "PRIVACY_MIGRATION_REQUIRED", "approved_for_analysis": False,
                    "manifest": str(path), **requirement}
    deterministic = manifest.get("deterministic_validation")
    original_deletion = manifest.get("original_deletion")
    approval_contract_valid = (
        manifest.get("status") in {"PASS", "SANITIZED"}
        and isinstance(deterministic, dict)
        and deterministic.get("passed") is True
        and isinstance(original_deletion, dict)
        and original_deletion.get("succeeded") is True
    )
    approved = manifest.get("approved_for_analysis") is True and approval_contract_valid
    if manifest.get("status") == "PURGED":
        state = "PURGED"
    elif manifest.get("status") in {"BLOCKED", "PURGE_PARTIAL_FAILURE"}:
        state = "PRIVACY_BLOCKED"
    elif manifest.get("approved_for_analysis") is True and not approval_contract_valid:
        state = "PRIVACY_BLOCKED"
    else:
        state = "PRIVACY_CLEARED" if approved else "AWAITING_PRIVACY_REVIEW"
    # A subsequent raw drop never inherits clearance from the previous batch.
    # Inspect directory entries only; Codex remains the first semantic reader.
    incoming = requirement["case_root"] / "incoming"
    if approved and incoming.exists() and any(incoming.iterdir()):
        state, approved = "AWAITING_PRIVACY_REVIEW", False
    return {"state": state, "approved_for_analysis": approved, "manifest": str(path), **requirement}


def assert_case_privacy_cleared(case_directory: str | Path) -> None:
    status = inspect_privacy_status(case_directory)
    if status["state"] == "NOT_REQUIRED":
        return
    if status["state"] != "PRIVACY_CLEARED" or not status["approved_for_analysis"]:
        raise ValueError(f"PRIVACY GATE: analyse interdite; état {status['state']} et approved_for_analysis != true.")
    from .contract_policy import assert_contract_permission
    assert_contract_permission(case_directory)


def _find_case_root(path: str | Path) -> Path | None:
    current = Path(path).expanduser().resolve(strict=False)
    if current.is_file() or current.suffix:
        current = current.parent
    for candidate in (current, *current.parents):
        if (candidate / "workspace.json").is_file() or (candidate / "case_manifest.json").is_file():
            return candidate
        if all((candidate / name).is_dir() for name in ("derived", "evidence", "investigation")):
            return candidate
    return _conventional_case_root(current)


def case_root_for_path(path: str | Path) -> Path | None:
    """Return the owning recognized case without reading client content."""

    return _find_case_root(path)


def assert_source_approved_for_analysis(source: str | Path, *, output_directory: str | Path | None = None) -> None:
    source_root = _find_case_root(source)
    output_root = _find_case_root(output_directory) if output_directory else None
    roots = [root for root in (source_root, output_root) if root is not None]
    if not roots:
        return
    root = roots[0]
    if any(item.resolve() != root.resolve() for item in roots[1:]):
        raise ValueError("La source et la sortie appartiennent à des workspaces clients différents.")
    status = inspect_privacy_status(root)
    assert_case_privacy_cleared(root)
    if status["state"] == "NOT_REQUIRED":
        return
    if source_root is not None and output_directory is not None and output_root is None:
        raise ValueError("PRIVACY GATE: les sorties d'un dossier réel doivent rester dans le même workspace isolé.")
    source_path = Path(source).resolve()
    allowed = [root / "sanitized", root / "normalized", root / "derived", root / "processed"]
    if not any(_within(source_path, item) for item in allowed):
        raise ValueError("PRIVACY GATE: la source analytique doit provenir de sanitized/ ou d'un dérivé autorisé après clearance.")
    if output_directory is not None:
        layout = resolve_case_layout(root)["layout"]
        output_path = Path(output_directory).resolve(strict=False)
        allowed_outputs = (
            [root / "processed", root / "scratch", root / "outputs"]
            if layout == "STANDARD_WORKSPACE"
            else [root / name for name in ("normalized", "derived", "evidence", "investigation", "scratch", "outputs")]
        )
        if not any(_within(output_path, directory) for directory in allowed_outputs):
            raise ValueError("PRIVACY GATE: destination analytique hors des zones isolées autorisées.")
    if _within(source_path, root / "sanitized") and source_path.is_file():
        manifest = _read_json(privacy_manifest_path(root))
        hashes = {item.get("sanitized_sha256") for item in manifest.get("files", []) if isinstance(item, dict)}
        if _sha256(source_path) not in hashes:
            raise ValueError("PRIVACY GATE: le hash de la source sanitized n'est pas approuvé par le manifest.")


def stage_incoming_drop(source_directory: str | Path, case_directory: str | Path) -> dict[str, Any]:
    """Copy a drop to incoming without parsing or semantically classifying it."""

    source = Path(source_directory).resolve()
    requirement = privacy_requirement(case_directory)
    if requirement["legacy"]:
        raise ValueError("Workspace historique: migration privacy explicite requise.")
    if not requirement["required"]:
        raise ValueError("stage_incoming_drop est réservé aux dossiers réels soumis au privacy gate.")
    root: Path = requirement["case_root"]
    if inspect_privacy_status(root)["state"] == "PURGED":
        raise ValueError("Dossier PURGED: aucun nouveau staging ou traitement n'est autorisé.")
    from .contract_policy import assert_contract_permission, contract_policy_digest, POLICY_PATH
    contract = assert_contract_permission(root, operation="staging")
    incoming = root / "incoming"
    incoming.mkdir(parents=True, exist_ok=True)
    if not source.is_dir():
        raise ValueError(f"Dossier reçu introuvable: {source}.")
    if any(incoming.iterdir()):
        raise FileExistsError("incoming/ contient déjà un dépôt; aucun écrasement silencieux.")
    files = _safe_files(source)
    if not files:
        raise ValueError("Le dépôt reçu ne contient aucun fichier.")
    copied: list[str] = []
    for path in files:
        relative = path.relative_to(source)
        target = incoming / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(path, target)
        copied.append(str(relative))
    receipt = {
        "schema_version": "indicia-incoming-receipt-v1",
        "contract_policy_ref": POLICY_PATH,
        "contract_policy_sha256": contract_policy_digest(contract),
        "workspace_id": root.name,
        "received_at_utc": _now(),
        "file_count": len(copied),
        "logical_paths": copied,
        "content_semantically_inspected": False,
        "next_action": "CODEX_PRIVACY_GATE",
    }
    _atomic_json(root / "privacy" / "incoming_receipt.json", receipt)
    return receipt


@dataclass
class _Table:
    name: str
    headers: list[str]
    rows: list[list[Any]]


def _csv_table(path: Path) -> _Table:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                sample = handle.read(8192)
                handle.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                data = [list(row) for row in csv.reader(handle, dialect)]
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError(f"Encodage CSV non pris en charge: {path.name}.")
    if not data:
        return _Table("CSV", [], [])
    width = max(len(row) for row in data)
    headers = [str(value or "") for value in data[0]] + [""] * (width - len(data[0]))
    return _Table("CSV", headers, [row + [""] * (width - len(row)) for row in data[1:]])


def _xlsx_tables(path: Path) -> list[_Table]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ValueError("openpyxl est requis pour vérifier un XLSX.") from exc
    # Formula text is part of the industrial source and must not be silently
    # changed during sanitation, so compare formulas rather than cached values.
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        result: list[_Table] = []
        for sheet in workbook.worksheets:
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            if not rows:
                result.append(_Table(sheet.title, [], []))
                continue
            width = max(len(row) for row in rows)
            headers = [str(value or "") for value in rows[0]] + [""] * (width - len(rows[0]))
            result.append(_Table(sheet.title, headers, [row + [None] * (width - len(row)) for row in rows[1:]]))
        return result
    finally:
        workbook.close()


def _tables(path: Path) -> list[_Table]:
    if path.suffix.casefold() == ".csv":
        return [_csv_table(path)]
    if path.suffix.casefold() == ".xlsx":
        return _xlsx_tables(path)
    return []


def _flatten_text(path: Path, tables: list[_Table]) -> str:
    if tables:
        return "\n".join(str(value or "") for table in tables for row in [table.headers, *table.rows] for value in row)
    if path.suffix.casefold() in _TEXT_EXTENSIONS:
        for encoding in ("utf-8-sig", "utf-8", "latin-1"):
            try:
                return path.read_text(encoding=encoding)
            except UnicodeDecodeError:
                continue
    raise ValueError(f"Format impossible à vérifier avec confiance: {path.suffix or 'sans extension'}.")


def _header_kind(header: str) -> str:
    slug = _slug(header)
    parts = set(slug.split("_"))
    if slug in _FREE_TEXT_HEADER_TOKENS:
        return "free_text"
    if slug in _PERSONAL_HEADER_TOKENS or any(token in slug for token in ("email", "telephone", "phone", "matricule")):
        return "relation" if slug in _RELATION_HEADER_TOKENS else "personal"
    if parts & current_preservation_policy().header_parts:
        return "industrial"
    return "other"


def _scan_text_patterns(
    text: str,
    categories: dict[str, int],
    sensitive_values: set[str],
) -> None:
    emails = _EMAIL.findall(text)
    phones = _PHONE.findall(text)
    secrets = [match.group(0) for pattern in _SECRET_PATTERNS for match in pattern.finditer(text)]
    if emails:
        categories["EMAIL"] = categories.get("EMAIL", 0) + len(emails)
        sensitive_values.update(emails)
        # Consumer mailbox domains can also be legitimate B2B contacts. Only
        # an explicit personal/private label or the reviewed category upgrades
        # risk; a provider domain alone is not a semantic determination.
        if _PERSONAL_EMAIL_LABEL.search(text):
            categories["PERSONAL_EMAIL"] = categories.get("PERSONAL_EMAIL", 0) + len(emails)
    if phones:
        categories["PHONE"] = categories.get("PHONE", 0) + len(phones)
        sensitive_values.update(phones)
    if secrets:
        categories["AUTHENTICATION_SECRET"] = categories.get("AUTHENTICATION_SECRET", 0) + len(secrets)
        sensitive_values.update(secrets)
    medical = _MEDICAL.findall(text)
    hr_sensitive = _HR_SENSITIVE.findall(text)
    if medical:
        categories["MEDICAL_DATA"] = categories.get("MEDICAL_DATA", 0) + len(medical)
    if hr_sensitive:
        categories["HR_SENSITIVE"] = categories.get("HR_SENSITIVE", 0) + len(hr_sensitive)
    for name, pattern in (("IDENTITY_DOCUMENT", _IDENTITY_DOCUMENT),
                          ("RESIDENTIAL_ADDRESS", _RESIDENTIAL_ADDRESS)):
        found = pattern.findall(text)
        if found:
            categories[name] = categories.get(name, 0) + len(found)


def _scan(path: Path) -> dict[str, Any]:
    tables = _tables(path)
    categories: dict[str, int] = {}
    sensitive_values: set[str] = set()
    inspection = inspect_document(path)
    if path.suffix.casefold() == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))

        def inspect_json(value, header=""):
            if isinstance(value, dict):
                for key, item in value.items():
                    inspect_json(item, key)
            elif isinstance(value, list):
                for item in value:
                    inspect_json(item, header)
            elif value is not None:
                text = str(value).strip()
                kind = _header_kind(header)
                digest = header.endswith("sha256") and re.fullmatch(r"[a-f0-9]{64}", text)
                if kind != "industrial" and not digest:
                    _scan_text_patterns(text, categories, sensitive_values)
                if text and kind in {"personal", "relation"}:
                    if kind == "relation" and _PSEUDONYM.fullmatch(text):
                        return
                    category = "INDIVIDUAL_IDENTIFIER" if kind == "relation" else "EXPLICIT_PERSONAL_COLUMN"
                    categories[category] = categories.get(category, 0) + 1
                    sensitive_values.add(text)
        inspect_json(payload)
    elif inspection is not None:
        _scan_text_patterns(inspection.native_text, categories, sensitive_values)
    elif not tables:
        _scan_text_patterns(_flatten_text(path, tables), categories, sensitive_values)
    for table in tables:
        for index, header in enumerate(table.headers):
            kind = _header_kind(header)
            nonempty = [str(row[index]).strip() for row in table.rows if index < len(row) and str(row[index] or "").strip()]
            # Industrial identifiers stay intact even when their spelling looks
            # like a person or a phone number. Codex has already reviewed their
            # meaning; generic post-check patterns only inspect other fields.
            if kind != "industrial":
                _scan_text_patterns("\n".join(nonempty), categories, sensitive_values)
            if kind not in {"personal", "relation"}:
                continue
            category = "INDIVIDUAL_IDENTIFIER" if kind == "relation" else "EXPLICIT_PERSONAL_COLUMN"
            if kind == "relation" and nonempty and all(_PSEUDONYM.fullmatch(value) for value in nonempty):
                continue
            if nonempty:
                categories[category] = categories.get(category, 0) + len(nonempty)
                sensitive_values.update(value for value in nonempty if len(value) >= 3)
    return {"tables": tables, "categories": categories, "sensitive_values": sensitive_values,
            "inspection": inspection.audit() if inspection is not None else None,
            "inspection_text": inspection.native_text if inspection is not None else None}


def _assert_pdf_redaction(original_scan: dict[str, Any], residual_scan: dict[str, Any], transforms: list) -> None:
    """A redacted value cannot survive in page text or PDF metadata without its label."""
    before, after = original_scan["inspection_text"], residual_scan["inspection_text"]
    if before is None or after is None:
        return
    removed = {item.get("category") for item in transforms if isinstance(item, dict)
               and item.get("action") == "REMOVED"}
    forbidden: set[str] = set()
    if "EMAIL" in removed:
        forbidden.update(_EMAIL.findall(before))
    if "PHONE" in removed:
        forbidden.update(_PHONE.findall(before))
    if "PERSONAL_EMAIL" in removed:
        for label in _PERSONAL_EMAIL_LABEL.finditer(before):
            email = _EMAIL.search(before, label.end(), min(len(before), label.end() + 120))
            if email is not None:
                forbidden.add(email.group(0))
    if "AUTHENTICATION_SECRET" in removed:
        for pattern in _SECRET_PATTERNS:
            for match in pattern.finditer(before):
                value = match.group(0)
                forbidden.add(value)
                if "=" in value:
                    forbidden.add(value.rsplit("=", 1)[-1])
                if ":" in value:
                    forbidden.add(value.rsplit(":", 1)[-1])
    if any(len(value) >= 6 and value.casefold() in after.casefold() for value in forbidden):
        raise ValueError("SANITIZATION_FAILED: removed value survives in PDF content or metadata.")


def _assess_inspection(spec: dict[str, Any], raw_scan: dict[str, Any]) -> dict[str, Any]:
    """Require accountable, hash-bound human review for every visual component."""
    audit = raw_scan["inspection"]
    reviews = spec.get("visual_reviews", [])
    if not isinstance(reviews, list):
        raise ValueError("Visual privacy reviews must be a list.")
    reviewed_categories: dict[str, int] = {}
    if audit is None:
        if reviews:
            raise ValueError("Visual review cannot be attached to a nonvisual source.")
    else:
        limitations = set(audit["limitations"])
        accountably_visual = {"PDF_FORM_FIELDS_REQUIRE_REVIEW", "PDF_ANNOTATIONS_REQUIRE_REVIEW"}
        if limitations - accountably_visual:
            raise ValueError("UNINSPECTABLE_COMPONENT: source contains unresolved document components.")
        required = set(audit["visual_components"])
        if required and not current_preservation_policy().risk.visual_human_review:
            raise ValueError("UNINSPECTABLE_COMPONENT: this privacy profile has no visual review path.")
        if len(reviews) != len(required) or {r.get("location") for r in reviews if isinstance(r, dict)} != required:
            raise ValueError("PARTIAL_DOCUMENT_INSPECTION: every visual component needs one review.")
        for review in reviews:
            if not isinstance(review, dict) or set(review) != {
                "location", "source_sha256", "inspection_version", "reviewer_role", "reviewed_at_utc",
                "decision", "detected_categories", "business_evidence_preserved", "prompt_injection_ignored",
            }:
                raise ValueError("Invalid closed visual-inspection contract.")
            if (review["source_sha256"] != audit["source_sha256"]
                    or review["inspection_version"] != INSPECTION_VERSION
                    or review["reviewer_role"] != "HUMAN"
                    or review["decision"] != "PASS"
                    or review["business_evidence_preserved"] is not True
                    or review["prompt_injection_ignored"] is not True):
                raise ValueError("MODEL_INSPECTION_FAILED: source-bound human visual approval required.")
            try:
                timestamp = datetime.fromisoformat(str(review["reviewed_at_utc"]).replace("Z", "+00:00"))
            except ValueError as exc:
                raise ValueError("Visual review timestamp invalid.") from exc
            if timestamp.tzinfo is None:
                raise ValueError("Visual review timestamp needs timezone.")
            categories = review["detected_categories"]
            if (not isinstance(categories, list) or len(categories) != len(set(categories))
                    or any(not _SAFE_CODE.fullmatch(str(category)) for category in categories)):
                raise ValueError("Visual review categories must be unique structured codes.")
            for category in categories:
                reviewed_categories[category] = reviewed_categories.get(category, 0) + 1
        if not audit["native_units"] and not required:
            raise ValueError("UNINSPECTABLE_COMPONENT: no readable or reviewed component.")
    for category in _clean_category_counts(spec.get("categories", [])):
        reviewed_categories[category["category"]] = max(reviewed_categories.get(category["category"], 0), category["count"])
    for category, count in reviewed_categories.items():
        raw_scan["categories"][category] = max(raw_scan["categories"].get(category, 0), count)
    categories = set(raw_scan["categories"])
    hard = categories & _HARD_PRIVACY_CATEGORIES
    risk_level = "HIGH" if hard or "AUTHENTICATION_SECRET" in categories else (
        "MODERATE" if categories - current_preservation_policy().risk.pass_categories else
        "LOW" if categories else "NONE")
    return {"inspectability": "INSPECTED" if audit is None or not audit["visual_components"] else "HUMAN_VISUAL_REVIEWED",
            "inspection": audit, "risk_level": risk_level,
            "detected_categories": sorted(categories),
            "human_visual_reviews": [{"location": r["location"], "reviewer_role": "HUMAN",
                                      "reviewed_at_utc": r["reviewed_at_utc"],
                                      "review_sha256": hashlib.sha256(json.dumps(r, sort_keys=True).encode()).hexdigest()}
                                     for r in reviews]}


def _same_value(left: Any, right: Any) -> bool:
    if left is None and right in (None, ""):
        return True
    if right is None and left in (None, ""):
        return True
    return left == right or str(left) == str(right)


def _identity_key(value: str) -> str:
    """Canonicalize an identity only for an ephemeral, in-process comparison."""

    return " ".join(unicodedata.normalize("NFKC", value).split()).casefold()


def _industrial_text_markers(path: Path) -> Counter[str]:
    inspection = inspect_document(path)
    if inspection is not None:
        if inspection.visual_components or inspection.limitations:
            raise ValueError("BUSINESS_EVIDENCE_LOSS: visual/opaque source cannot be safely rewritten.")
        text = inspection.native_text
    else:
        text = _flatten_text(path, [])
    markers: list[str] = []
    for pattern in current_preservation_policy().text_patterns:
        for match in pattern.finditer(text):
            marker = match.group(1) if match.lastindex else match.group(0)
            markers.append(marker.casefold())
    return Counter(markers)


def _compare_preservation(
    original: Path,
    sanitized: Path,
    *,
    require_relation_preservation: bool,
    identity_to_pseudonym: dict[str, str],
    pseudonym_to_identity: dict[str, str],
) -> dict[str, Any]:
    if original.suffix.casefold() != sanitized.suffix.casefold():
        raise ValueError("La sanitation ne peut pas changer le type de fichier tabulaire.")
    if original.suffix.casefold() == ".json":
        before = json.loads(original.read_text(encoding="utf-8"))
        after = json.loads(sanitized.read_text(encoding="utf-8"))

        def compare(left, right):
            if isinstance(left, dict):
                if not isinstance(right, dict) or set(right) - set(left):
                    raise ValueError("JSON structure changed during privacy cleanup.")
                for key, value in left.items():
                    kind = _header_kind(key)
                    if kind == "personal":
                        if key in right and right[key] not in (None, "", []):
                            raise ValueError("Personal JSON field must be removed or empty.")
                    elif kind == "relation":
                        new = right.get(key)
                        if not value and not new:
                            continue
                        if new is None and not require_relation_preservation:
                            continue
                        if not isinstance(value, str) or not isinstance(new, str) or not _PSEUDONYM.fullmatch(new):
                            raise ValueError("Stable JSON relation pseudonym required.")
                        identity, pseudonym = _identity_key(value), new.casefold()
                        if (identity in identity_to_pseudonym and identity_to_pseudonym[identity] != pseudonym
                                or pseudonym in pseudonym_to_identity and pseudonym_to_identity[pseudonym] != identity):
                            raise ValueError("Inconsistent or colliding JSON pseudonyms.")
                        identity_to_pseudonym[identity] = pseudonym
                        pseudonym_to_identity[pseudonym] = identity
                    else:
                        if key not in right:
                            raise ValueError("Business JSON field removed during privacy cleanup.")
                        compare(value, right[key])
            elif isinstance(left, list):
                if not isinstance(right, list) or len(left) != len(right):
                    raise ValueError("JSON rows changed during privacy cleanup.")
                for old, new in zip(left, right):
                    compare(old, new)
            elif type(left) is not type(right) or left != right:
                raise ValueError("Business JSON value changed during privacy cleanup.")

        compare(before, after)
        return {"tabular": False, "industrial_columns_checked": 0, "rows_checked": 0,
                "structured_json_values_preserved": True}
    if original.suffix.casefold() not in _TABULAR_EXTENSIONS:
        before_markers = _industrial_text_markers(original)
        after_markers = _industrial_text_markers(sanitized)
        missing = before_markers - after_markers
        if missing:
            raise ValueError("Un marqueur industriel textuel a été supprimé ou modifié.")
        return {
            "tabular": False,
            "industrial_columns_checked": 0,
            "rows_checked": 0,
            "industrial_text_markers_checked": sum(before_markers.values()),
        }
    before = _tables(original)
    after = _tables(sanitized)
    if [table.name for table in before] != [table.name for table in after]:
        raise ValueError("Les feuilles XLSX ou leur ordre ont changé pendant le privacy cleanup.")
    checked = rows_checked = 0
    for source, target in zip(before, after):
        if len(source.rows) != len(target.rows):
            raise ValueError("Le nombre de lignes a changé pendant le privacy cleanup.")
        target_by_slug = {_slug(header): index for index, header in enumerate(target.headers)}
        if len(target_by_slug) != len(target.headers):
            raise ValueError("En-têtes sanitized vides ou dupliqués: comparaison sûre impossible.")
        allowed_target = {_slug(header) for header in source.headers}
        if set(target_by_slug) - allowed_target:
            raise ValueError("Le privacy cleanup a ajouté ou renommé une colonne non déclarée.")
        for source_index, header in enumerate(source.headers):
            slug = _slug(header)
            kind = _header_kind(header)
            target_index = target_by_slug.get(slug)
            if kind in {"industrial", "other"}:
                if target_index is None:
                    raise ValueError(f"Colonne industrielle ou contextuelle supprimée: {header}.")
                for row_before, row_after in zip(source.rows, target.rows):
                    if not _same_value(row_before[source_index], row_after[target_index]):
                        raise ValueError(f"Valeur industrielle/contextuelle modifiée dans {header}.")
                checked += kind == "industrial"
            elif target_index is not None:
                values_before = [str(row[source_index] or "").strip() for row in source.rows]
                values_after = [str(row[target_index] or "").strip() for row in target.rows]
                if kind == "personal" and slug in _DIRECT_REMOVE_HEADER_TOKENS and any(values_after):
                    raise ValueError(f"La colonne personnelle explicite {header} doit être supprimée ou vidée.")
                if kind == "relation":
                    for old, new in zip(values_before, values_after):
                        if not old and not new:
                            continue
                        if old and not new and require_relation_preservation:
                            raise ValueError(f"Relation industrielle perdue pendant la pseudonymisation dans {header}.")
                        if old and new:
                            if not _PSEUDONYM.fullmatch(new):
                                raise ValueError(f"Pseudonyme invalide dans {header}: format stable requis.")
                            identity = _identity_key(old)
                            pseudonym = new.casefold()
                            if identity in identity_to_pseudonym and identity_to_pseudonym[identity] != pseudonym:
                                raise ValueError("Pseudonymisation inter-fichiers instable.")
                            if pseudonym in pseudonym_to_identity and pseudonym_to_identity[pseudonym] != identity:
                                raise ValueError("Collision de pseudonymes inter-fichiers.")
                            identity_to_pseudonym[identity] = pseudonym
                            pseudonym_to_identity[pseudonym] = identity
            elif kind == "relation" and require_relation_preservation:
                raise ValueError(f"Colonne relationnelle supprimée pendant une pseudonymisation: {header}.")
        rows_checked += len(source.rows)
    return {"tabular": True, "industrial_columns_checked": checked, "rows_checked": rows_checked}


def _clean_category_counts(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        raise ValueError("detected_categories doit être une liste agrégée.")
    clean: list[dict[str, Any]] = []
    for item in value:
        if not isinstance(item, dict) or not _SAFE_CODE.fullmatch(str(item.get("category", ""))):
            raise ValueError("Catégorie privacy invalide.")
        if set(item) - {"category", "count", "action", "file_ids"}:
            raise ValueError("Une catégorie privacy contient des champs interdits; ne jamais recopier les valeurs.")
        count = item.get("count")
        if not isinstance(count, int) or isinstance(count, bool) or count < 0:
            raise ValueError("Le compte privacy doit être un entier positif ou nul.")
        action = str(item.get("action", ""))
        file_ids = item.get("file_ids", [])
        if not _SAFE_CODE.fullmatch(action) or not isinstance(file_ids, list) or any(not _FILE_ID.fullmatch(str(file_id)) for file_id in file_ids):
            raise ValueError("Action ou file_ids privacy invalides.")
        clean.append({"category": str(item["category"]), "count": count, "action": action, "file_ids": list(file_ids)})
    return clean


def _assert_no_pre_gate_derivatives(root: Path, layout: str) -> None:
    """Prevent clearance when client content already escaped the gate zones."""

    if layout == "STANDARD_WORKSPACE":
        checked = (root / "processed", root / "scratch", root / "outputs", root / "retained_derived")
        allowed: set[Path] = {
            root / "processed" / "investigation_state.json",
            root / "processed" / "questions.json",
        }
    else:
        checked = (
            root / "normalized", root / "derived", root / "evidence", root / "scratch",
            root / "investigation", root / "outputs", root / "retained_derived",
        )
        allowed = {
            root / "investigation" / "case_state.json",
            root / "investigation" / "investigation_state.json",
            root / "investigation" / "questions.json",
        }
    from .artifact_store import LOCK
    analysis = root / ("processed" if layout == "STANDARD_WORKSPACE" else "investigation")
    lock = analysis / LOCK
    if lock.is_file() and not lock.is_symlink() and lock.stat().st_size == 0:
        allowed.add(lock)  # Empty synchronization metadata, never a content-bearing journal.
    escaped = [path for directory in checked for path in _safe_files(directory) if path not in allowed]
    if escaped:
        raise ValueError(
            "PRIVACY GATE: artefacts client présents hors incoming/privacy avant clearance; "
            "purge manuelle et nouveau gate requis."
        )


def _assert_privacy_auxiliaries_bounded(root: Path, review_file: Path) -> None:
    privacy_root = root / "privacy"
    candidate_root = privacy_root / "candidate"
    allowed = {
        privacy_root / "CODEX_PRIVACY_REVIEW_TEMPLATE.json",
        privacy_root / "RETENTION_POLICY_TEMPLATE.json",
        privacy_root / "incoming_receipt.json",
        privacy_root / "privacy_manifest.json",
        privacy_root / "retention_policy.json",
        review_file,
    }
    unexpected = [
        path for path in _safe_files(privacy_root)
        if not _within(path, candidate_root) and path not in allowed
    ]
    if unexpected:
        raise ValueError("PRIVACY GATE: artefact auxiliaire inattendu hors privacy/candidate/.")


def _manifest_base(root: Path, review: dict[str, Any], status: str) -> dict[str, Any]:
    extraction = _read_json(root / "contracts/contract_extraction.json")
    return {
        "contract_authorization_ref": extraction.get("authorization_ref"),
        "schema_version": MANIFEST_SCHEMA,
        "workspace_id": root.name,
        "received_at_utc": review.get("received_at_utc"),
        "validated_at_utc": _now(),
        "policy_version": POLICY_VERSION,
        "business_preservation_policy": current_preservation_policy().policy_id,
        "risk_policy_version": current_preservation_policy().risk.policy_id,
        "status": status,
        "detected_categories": _clean_category_counts(review.get("detected_categories", [])),
        "transformations": [],
        "files": [],
        "deterministic_validation": {"passed": False, "checks": [], "errors": []},
        "original_deletion": {
            "required": True,
            "attempted": False,
            "succeeded": False,
            "partial": False,
            "deleted_file_count": 0,
            "remaining_file_count": None,
            "deleted_at_utc": None,
            "errors": [],
        },
        "approved_for_analysis": False,
    }


def _blocked_temporary_cleanup(
    root: Path,
    manifest: dict[str, Any],
    *,
    delete_tree: Callable[[Path], None],
    review_file: Path | None = None,
    known_incoming_files: list[Path] | None = None,
    deletion_authorized: bool = True,
) -> None:
    """Best-effort fail-closed cleanup with a content-free, verified audit."""

    incoming = root / "incoming"
    candidate_root = root / "privacy" / "candidate"
    staging = root / "privacy" / ".promotion_staging"
    deletion = manifest["original_deletion"]
    if not deletion_authorized:
        deletion["errors"] = ["INCOMING_DELETE_NOT_SAFE:AUDIT_INVENTORY_INCOMPLETE"]
        manifest["temporary_cleanup"] = {
            "succeeded": False,
            "errors": ["RAW_CLEANUP_NOT_ATTEMPTED_INCOMPLETE_AUDIT"],
        }
        return
    try:
        originals = list(known_incoming_files) if known_incoming_files is not None else _safe_files(incoming)
    except (OSError, ValueError) as exc:
        deletion["errors"] = [f"INCOMING_INVENTORY_UNSAFE:{type(exc).__name__}"]
        manifest["temporary_cleanup"] = {
            "succeeded": False,
            "errors": ["RAW_CLEANUP_NOT_ATTEMPTED_UNSAFE_INVENTORY"],
        }
        return

    deletion["attempted"] = True
    incoming_error: str | None = None
    if incoming.exists():
        try:
            delete_tree(incoming)
        except OSError as exc:
            incoming_error = f"INCOMING_DELETE_FAILED:{type(exc).__name__}"

    remaining_count: int | None
    try:
        remaining_count = len(_safe_files(incoming))
    except (OSError, ValueError) as exc:
        remaining_count = None
        if incoming_error is None:
            incoming_error = f"INCOMING_DELETE_VERIFY_FAILED:{type(exc).__name__}"
    deleted_count = (
        max(0, len(originals) - remaining_count)
        if remaining_count is not None else 0
    )
    deletion.update({
        "succeeded": incoming_error is None and remaining_count == 0,
        "partial": bool(
            incoming_error is not None
            and remaining_count is not None
            and 0 < deleted_count < len(originals)
        ),
        "deleted_file_count": deleted_count,
        "remaining_file_count": remaining_count,
        "deleted_at_utc": _now() if incoming_error is None and remaining_count == 0 else None,
        "errors": [] if incoming_error is None else [incoming_error],
    })

    auxiliary_errors: list[str] = []
    for target, code in (
        (candidate_root, "CANDIDATE_DELETE_FAILED"),
        (staging, "PROMOTION_STAGING_DELETE_FAILED"),
    ):
        if target.exists():
            try:
                delete_tree(target)
            except OSError as exc:
                auxiliary_errors.append(f"{code}:{type(exc).__name__}")
    allowed_review = (
        review_file is not None
        and review_file.parent == (root / "privacy").resolve()
        and review_file.name not in {
            "privacy_manifest.json",
            "CODEX_PRIVACY_REVIEW_TEMPLATE.json",
            "RETENTION_POLICY_TEMPLATE.json",
            "retention_policy.json",
        }
    )
    files_to_unlink = [(root / "privacy" / "incoming_receipt.json", "INCOMING_RECEIPT_DELETE_FAILED")]
    if allowed_review:
        files_to_unlink.append((review_file, "REVIEW_DELETE_FAILED"))
    for target, code in files_to_unlink:
        if target.exists():
            try:
                target.unlink()
            except OSError as exc:
                auxiliary_errors.append(f"{code}:{type(exc).__name__}")
    manifest["temporary_cleanup"] = {
        "succeeded": deletion["succeeded"] and not auxiliary_errors,
        "errors": auxiliary_errors,
    }


def _merge_detected_category(manifest: dict[str, Any], category: str, count: int, file_id: str) -> None:
    for item in manifest["detected_categories"]:
        if item["category"] == category:
            item["count"] = max(int(item["count"]), int(count))
            item["file_ids"] = list(dict.fromkeys([*item.get("file_ids", []), file_id]))
            return
    manifest["detected_categories"].append({
        "category": category,
        "count": int(count),
        "action": "DETERMINISTIC_POST_CHECK",
        "file_ids": [file_id],
    })


def _sync_lifecycle(root: Path, state: str, *, reason: str) -> None:
    analysis = Path(resolve_case_layout(root)["analysis_root"])
    path = analysis / "investigation_state.json"
    if path.is_file():
        payload = _read_json(path)
        life = payload.get("client_lifecycle", {})
        if state in {"PRIVACY_CLEARED", "PRIVACY_BLOCKED"} and life.get("state") in {"ANALYZING", "WAITING_FOR_REQUIRED_INFORMATION", "RESUMING", "FINALIZABLE", "DELIVERABLE"}:
            # Privacy is a gate over an existing investigation, never permission to
            # erase its pending questions, answer history or resume obligation.
            life["privacy_approved_for_analysis"] = state == "PRIVACY_CLEARED"
            life["history"].append({"at_utc": _now(), "action": "additional_privacy_review", "privacy_state": state, "reason": reason})
            if state == "PRIVACY_CLEARED":
                life["resume_required"] = True
                if life["state"] != "WAITING_FOR_REQUIRED_INFORMATION":
                    life["state"] = "RESUMING"
            elif life["state"] == "DELIVERABLE":
                life["state"] = "FINALIZABLE"
            _atomic_json(path, payload)
            return
    from .client_lifecycle import synchronize_privacy_state
    synchronize_privacy_state(root, state, reason=reason)

def _update_case_privacy_summary(root: Path, *, approved: bool, status: str) -> None:
    for name in ("workspace.json", "case_manifest.json"):
        path = root / name
        if not path.is_file():
            continue
        payload = _read_json(path)
        privacy = payload.setdefault("privacy", {})
        privacy["approved_for_analysis"] = approved
        privacy["last_gate_status"] = status
        privacy["manifest"] = "privacy/privacy_manifest.json"
        if status == "PURGED":
            payload["status"] = "purged"
        elif status == "PURGE_PARTIAL_FAILURE":
            payload["status"] = "purge_partial_failure"
        else:
            payload["status"] = "privacy_cleared" if approved else "privacy_blocked"
        _atomic_json(path, payload)
        return


def _validate_codex_privacy_review_impl(
    case_directory: str | Path,
    review_path: str | Path,
    *,
    delete_tree: Callable[[Path], None] = shutil.rmtree,
    prior_manifest: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Fail-closed post-check and promotion after Codex's semantic review."""

    requirement = privacy_requirement(case_directory)
    if requirement["legacy"]:
        raise ValueError("Workspace historique: exécuter une migration privacy explicite; aucune migration automatique.")
    if not requirement["required"]:
        raise ValueError("Le dossier est explicitement synthétique et ne requiert pas ce gate.")
    root: Path = requirement["case_root"]
    if inspect_privacy_status(root)["state"] == "PURGED":
        raise ValueError("Dossier PURGED: le privacy gate ne peut pas être rouvert.")
    incoming = root / "incoming"
    candidate_root = root / "privacy" / "candidate"
    sanitized_root = root / "sanitized"
    review_input = Path(review_path)
    if review_input.is_symlink():
        raise ValueError("La privacy review ne peut pas être un lien symbolique.")
    review_file = review_input.resolve()
    if review_file.parent != (root / "privacy").resolve():
        raise ValueError("La privacy review doit être un fichier direct de privacy/ dans le workspace.")
    review = _read_json(review_file)
    allowed_review_fields = {
        "schema_version", "policy_version", "workspace_id", "received_at_utc", "status",
        "codex_semantic_review", "detected_categories", "files", "blocked_reasons",
    }
    if set(review) - allowed_review_fields:
        raise ValueError("La privacy review contient des champs libres interdits.")
    if review.get("schema_version") != REVIEW_SCHEMA or review.get("policy_version") != POLICY_VERSION:
        raise ValueError("Schéma ou policy de privacy review invalide.")
    if review.get("workspace_id") != root.name:
        raise ValueError("workspace_id de la privacy review incohérent.")
    semantic = review.get("codex_semantic_review")
    if (
        not isinstance(semantic, dict)
        or set(semantic) != {"completed", "first_substantive_reader_attested"}
        or semantic.get("completed") is not True
        or semantic.get("first_substantive_reader_attested") is not True
    ):
        raise ValueError("Codex doit attester la revue sémantique initiale du contenu brut.")
    try:
        received_at = datetime.fromisoformat(str(review.get("received_at_utc", "")).replace("Z", "+00:00"))
    except ValueError as exc:
        raise ValueError("received_at_utc doit être un timestamp ISO avec fuseau.") from exc
    if received_at.tzinfo is None:
        raise ValueError("received_at_utc doit inclure un fuseau.")
    status = review.get("status")
    if status not in {"PASS", "SANITIZED", "BLOCKED"}:
        raise ValueError("Statut privacy inconnu.")
    manifest = _manifest_base(root, review, status)
    if prior_manifest is not None:
        manifest["previous_reviews"] = [*prior_manifest.get("previous_reviews", []),
            {k: v for k, v in prior_manifest.items() if k != "previous_reviews"}]
        manifest["supplemental_review"] = True
        manifest["files"] = list(prior_manifest["files"])
    blocked_reasons = review.get("blocked_reasons", [])
    if not isinstance(blocked_reasons, list) or any(not _SAFE_CODE.fullmatch(str(reason)) for reason in blocked_reasons):
        raise ValueError("blocked_reasons doit contenir uniquement des codes structurés sans donnée source.")
    manifest["blocked_reason_codes"] = list(blocked_reasons)
    incoming_files = _safe_files(incoming)
    if not incoming_files:
        raise ValueError("incoming/ est vide: aucun brut temporaire à valider.")
    if prior_manifest is None:
        _assert_no_pre_gate_derivatives(root, requirement["layout"])
    _assert_privacy_auxiliaries_bounded(root, review_file)
    file_specs = review.get("files")
    if not isinstance(file_specs, list) or len(file_specs) != len(incoming_files):
        raise ValueError("La privacy review doit couvrir exactement tous les fichiers incoming.")
    by_source: dict[str, dict[str, Any]] = {}
    file_ids: set[str] = {item["file_id"] for item in prior_manifest["files"]} if prior_manifest else set()
    for spec in file_specs:
        if not isinstance(spec, dict) or set(spec) - {"file_id", "source", "action", "sanitized", "categories", "transformations", "industrial_content_absent", "visual_reviews", "business_confidentiality"}:
            raise ValueError("Contrat de fichier privacy invalide ou champ libre interdit.")
        source = str(spec.get("source", ""))
        file_id = str(spec.get("file_id", ""))
        if not _FILE_ID.fullmatch(file_id) or file_id in file_ids:
            raise ValueError("file_id privacy absent, invalide ou dupliqué.")
        file_ids.add(file_id)
        _clean_category_counts(spec.get("categories", []))
        if spec.get("business_confidentiality", "RESTRICTED_CLIENT") not in {
            "PUBLIC", "BUSINESS_CONFIDENTIAL", "RESTRICTED_CLIENT"
        }:
            raise ValueError("Business confidentiality must use a closed category.")
        if source in by_source or not source:
            raise ValueError("Source privacy absente ou dupliquée.")
        by_source[source] = spec
    expected_sources = {str(path.relative_to(root)) for path in incoming_files}
    if set(by_source) != expected_sources:
        raise ValueError("La liste des sources review ne correspond pas exactement à incoming/.")
    if status == "BLOCKED":
        if not blocked_reasons:
            raise ValueError("BLOCKED exige au moins une raison structurée.")
        if not any(spec.get("action") == "BLOCKED" for spec in file_specs) or any(
            spec.get("action") not in {"PASS", "SANITIZED", "REMOVE", "BLOCKED"}
            for spec in file_specs
        ):
            raise ValueError("BLOCKED exige au moins une action fichier BLOCKED et aucune action inconnue.")
        for path in incoming_files:
            spec = by_source[str(path.relative_to(root))]
            try:
                blocked_scan = _scan(path)
                blocked_assessment = _assess_inspection(spec, blocked_scan)
            except (ValueError, OSError) as exc:
                # A blocked source may be corrupt or intentionally left unread.
                # Preserve the failure class, never the source/error text.
                blocked_assessment = {"inspectability": "UNINSPECTABLE_OR_UNREVIEWED",
                                      "inspection": None, "risk_level": "UNKNOWN",
                                      "detected_categories": [], "inspection_failure": type(exc).__name__}
            manifest["files"].append({
                "file_id": spec["file_id"], "original_sha256": _sha256(path),
                "file_type": path.suffix.casefold().lstrip(".") or "unknown",
                "status": spec["action"], "sanitized_sha256": None,
                "privacy_assessment": {**blocked_assessment,
                    "business_confidentiality": spec.get("business_confidentiality", "RESTRICTED_CLIENT"),
                    "risk_policy_version": current_preservation_policy().risk.policy_id,
                    "required_action": spec["action"]},
            })
        manifest["deterministic_validation"]["checks"].append("blocked_without_analytical_promotion")
        manifest["deterministic_validation"]["passed"] = True
        _blocked_temporary_cleanup(
            root,
            manifest,
            delete_tree=delete_tree,
            review_file=review_file,
            known_incoming_files=incoming_files,
        )
        _atomic_json(privacy_manifest_path(root), manifest)
        _sync_lifecycle(root, "PRIVACY_BLOCKED", reason="codex_privacy_review_blocked")
        _update_case_privacy_summary(root, approved=False, status="BLOCKED")
        return manifest
    if status == "PASS" and any(spec.get("action") != "PASS" for spec in file_specs):
        raise ValueError("PASS exige une action PASS pour chaque fichier.")
    if status == "PASS" and {item["category"] for item in manifest["detected_categories"]} - current_preservation_policy().risk.pass_categories:
        raise ValueError("PASS est incompatible avec des catégories privacy détectées.")
    if status == "SANITIZED" and not any(spec.get("action") in {"SANITIZED", "REMOVE"} for spec in file_specs):
        raise ValueError("SANITIZED exige au moins une transformation ou suppression.")
    staging = root / "privacy" / ".promotion_staging"
    if staging.exists():
        delete_tree(staging)
    staging.mkdir(parents=True)
    all_raw_sensitive: set[str] = set()
    validation_checks: list[str] = []
    target_paths: set[str] = set()
    identity_to_pseudonym: dict[str, str] = {}
    pseudonym_to_identity: dict[str, str] = {}
    declared_transformations = {
        (item["category"], item["action"]): int(item["count"])
        for item in manifest["detected_categories"]
    }
    try:
        for index, original in enumerate(incoming_files, 1):
            source_key = str(original.relative_to(root))
            spec = by_source[source_key]
            file_id = str(spec["file_id"])
            action = spec.get("action")
            if action not in {"PASS", "SANITIZED", "REMOVE"}:
                raise ValueError(f"Action privacy invalide pour {file_id}.")
            raw_scan = _scan(original)
            assessment = _assess_inspection(spec, raw_scan)
            all_raw_sensitive.update(raw_scan["sensitive_values"])
            for category, count in raw_scan["categories"].items():
                _merge_detected_category(manifest, category, count, file_id)
            if set(raw_scan["categories"]) & _HARD_PRIVACY_CATEGORIES:
                raise ValueError(f"{file_id}: HIGH_RISK_PERSONAL_DATA; BLOCKED obligatoire.")
            entry = {
                "file_id": file_id, "original_sha256": _sha256(original),
                "file_type": original.suffix.casefold().lstrip(".") or "unknown",
                "status": action, "detected_categories": sorted(raw_scan["categories"]),
                "sanitized_sha256": None, "duplicate_of_file_id": None,
                "privacy_assessment": {**assessment,
                    "business_confidentiality": spec.get("business_confidentiality", "RESTRICTED_CLIENT"),
                    "risk_policy_version": current_preservation_policy().risk.policy_id,
                    "required_action": action},
            }
            if raw_scan["inspection"] is not None and entry["original_sha256"] != raw_scan["inspection"]["source_sha256"]:
                raise ValueError(f"{file_id}: SOURCE_CHANGED during privacy inspection.")
            if action == "REMOVE":
                if spec.get("industrial_content_absent") is not True:
                    raise ValueError(f"{file_id}: REMOVE exige l'attestation industrial_content_absent=true.")
                if raw_scan["inspection"] is not None and (
                    raw_scan["inspection"]["visual_components"] or _industrial_text_markers(original)
                ):
                    raise ValueError(f"{file_id}: BUSINESS_EVIDENCE_LOSS; document commercial ne peut être supprimé.")
                if any(_header_kind(header) == "industrial" for table in raw_scan["tables"] for header in table.headers):
                    raise ValueError(f"{file_id}: suppression interdite d'un fichier contenant des colonnes industrielles.")
                manifest["files"].append(entry)
                manifest["transformations"].append({"file_id": file_id, "category": "FILE_REMOVAL", "action": "REMOVED", "count": 1})
                continue
            if action == "PASS":
                if set(raw_scan["categories"]) - current_preservation_policy().risk.pass_categories:
                    raise ValueError(f"{file_id}: PASS refusé; motifs privacy déterministes détectés.")
                source_for_promotion = original
                relative_target = str(Path(source_key).relative_to("incoming"))
            else:
                if not isinstance(spec.get("transformations"), list) or not spec["transformations"]:
                    raise ValueError(f"{file_id}: SANITIZED exige des transformations structurées.")
                candidate = spec.get("sanitized")
                if not isinstance(candidate, str) or not candidate.startswith("privacy/candidate/"):
                    raise ValueError(f"{file_id}: candidat sanitized absent ou hors privacy/candidate/.")
                source_for_promotion = root / candidate
                if source_for_promotion.is_symlink() or not source_for_promotion.is_file() or not _within(source_for_promotion, candidate_root):
                    raise ValueError(f"{file_id}: candidat sanitized introuvable.")
                residual = _scan(source_for_promotion)
                _assess_inspection({"visual_reviews": [], "categories": []}, residual)
                if set(residual["categories"]) - current_preservation_policy().risk.pass_categories:
                    raise ValueError(f"{file_id}: post-check détecte encore {sorted(residual['categories'])}.")
                if original.suffix.casefold() == ".pdf":
                    _assert_pdf_redaction(raw_scan, residual, spec["transformations"])
                preservation = _compare_preservation(
                    original,
                    source_for_promotion,
                    require_relation_preservation=any(
                        transform.get("action") == "PSEUDONYMIZED"
                        for transform in spec.get("transformations", [])
                        if isinstance(transform, dict)
                    ),
                    identity_to_pseudonym=identity_to_pseudonym,
                    pseudonym_to_identity=pseudonym_to_identity,
                )
                validation_checks.append(f"{file_id}:industrial_preservation:{preservation['industrial_columns_checked']}")
                relative_target = str(Path(candidate).relative_to("privacy/candidate"))
                if Path(relative_target).is_absolute() or ".." in Path(relative_target).parts:
                    raise ValueError("Chemin sanitized non sûr.")
                for transform in spec.get("transformations", []):
                    if not isinstance(transform, dict) or set(transform) - {"category", "action", "count", "preserves_relations", "reason_code"}:
                        raise ValueError("Transformation privacy invalide; aucune valeur source n'est autorisée.")
                    if transform.get("preserves_relations") is not True:
                        raise ValueError("Toute transformation doit attester la préservation des relations industrielles utiles.")
                    if (
                        not _SAFE_CODE.fullmatch(str(transform.get("category", "")))
                        or str(transform.get("action", "")) not in {"REMOVED", "PSEUDONYMIZED"}
                        or not _SAFE_CODE.fullmatch(str(transform.get("reason_code", "")))
                        or not isinstance(transform.get("count"), int)
                        or isinstance(transform.get("count"), bool)
                        or transform.get("count") < 1
                    ):
                        raise ValueError("Transformation privacy non structurée ou compte invalide.")
                    declared_count = declared_transformations.get(
                        (str(transform["category"]), str(transform["action"])), 0
                    )
                    if declared_count < int(transform["count"]):
                        raise ValueError("Transformation absente ou sous-déclarée dans detected_categories.")
                    manifest["transformations"].append({"file_id": file_id, **transform})
            if relative_target in target_paths:
                raise ValueError("Deux fichiers privacy ciblent le même chemin sanitized.")
            target_paths.add(relative_target)
            target = staging / relative_target
            target.parent.mkdir(parents=True, exist_ok=True)
            if _sha256(original) != entry["original_sha256"]:
                raise ValueError(f"{file_id}: SOURCE_CHANGED before privacy promotion.")
            shutil.copy2(source_for_promotion, target)
            entry["sanitized_sha256"] = _sha256(target)
            if action == "PASS" and entry["sanitized_sha256"] != entry["original_sha256"]:
                raise ValueError(f"{file_id}: SOURCE_CHANGED during privacy promotion.")
            if action == "SANITIZED" and residual["inspection"] is not None and (
                entry["sanitized_sha256"] != residual["inspection"]["source_sha256"]
            ):
                raise ValueError(f"{file_id}: sanitized derivative changed during promotion.")
            manifest["files"].append(entry)
        review_serialized = json.dumps(review, ensure_ascii=False).casefold()
        leaked = [value for value in all_raw_sensitive if len(value) >= 4 and value.casefold() in review_serialized]
        if leaked:
            raise ValueError("La privacy review recopie une ou plusieurs valeurs supprimées; audit refusé.")
        hashes: dict[str, str] = {}
        for entry in manifest["files"]:
            digest = entry.get("original_sha256")
            if digest in hashes:
                entry["duplicate_of_file_id"] = hashes[digest]
            else:
                hashes[digest] = entry["file_id"]
        if prior_manifest is None:
            if sanitized_root.exists() and any(sanitized_root.iterdir()):
                raise FileExistsError("sanitized/ n'est pas vide; aucune promotion/écrasement silencieux.")
            sanitized_root.rmdir()
            os.replace(staging, sanitized_root)
        else:
            promoted = _safe_files(staging)
            if any((sanitized_root / item.relative_to(staging)).exists() for item in promoted):
                raise FileExistsError("Additional evidence cannot replace an existing sanitized document; use a new filename.")
            for item in promoted:
                target = sanitized_root / item.relative_to(staging)
                target.parent.mkdir(parents=True, exist_ok=True)
                os.replace(item, target)
            delete_tree(staging)
        cleanup_errors: list[str] = []
        for target in (incoming, candidate_root):
            if target.exists():
                try:
                    delete_tree(target)
                except OSError as exc:
                    cleanup_errors.append(f"{target.name}: {type(exc).__name__}")
        for target in (review_file, root / "privacy" / "incoming_receipt.json"):
            if target.exists():
                try:
                    target.unlink()
                except OSError as exc:
                    cleanup_errors.append(f"privacy_auxiliary: {type(exc).__name__}")
        if cleanup_errors:
            manifest["status"] = "BLOCKED"
            manifest["deterministic_validation"]["errors"] = cleanup_errors
            manifest["deterministic_validation"]["checks"] = validation_checks
            try:
                remaining_count = len(_safe_files(incoming))
            except (OSError, ValueError):
                remaining_count = None
            deleted_count = (
                max(0, len(incoming_files) - remaining_count)
                if remaining_count is not None else 0
            )
            manifest["original_deletion"] = {
                "required": True,
                "attempted": True,
                "succeeded": remaining_count == 0 and not any(
                    error.startswith("incoming:") for error in cleanup_errors
                ),
                "partial": bool(
                    remaining_count is not None
                    and 0 < deleted_count < len(incoming_files)
                ),
                "deleted_file_count": deleted_count,
                "remaining_file_count": remaining_count,
                "deleted_at_utc": _now() if remaining_count == 0 else None,
                "errors": [
                    "INCOMING_DELETE_FAILED"
                    for error in cleanup_errors if error.startswith("incoming:")
                ],
            }
            _atomic_json(privacy_manifest_path(root), manifest)
            _sync_lifecycle(root, "PRIVACY_BLOCKED", reason="temporary_original_deletion_failed")
            _update_case_privacy_summary(root, approved=False, status="BLOCKED")
            return manifest
        incoming.mkdir()
        candidate_root.mkdir(parents=True, exist_ok=True)
        manifest["deterministic_validation"] = {"passed": True, "checks": [
            "all_incoming_files_covered", "obvious_pii_and_secret_patterns_absent",
            "file_types_and_row_counts_preserved", "industrial_values_unchanged",
            "timestamps_energy_power_production_preserved", *validation_checks,
        ], "errors": []}
        manifest["original_deletion"] = {
            "required": True,
            "attempted": True,
            "succeeded": True,
            "partial": False,
            "deleted_file_count": len(incoming_files),
            "remaining_file_count": 0,
            "deleted_at_utc": _now(),
            "errors": [],
        }
        manifest["approved_for_analysis"] = True
        _atomic_json(privacy_manifest_path(root), manifest)
        _sync_lifecycle(root, "PRIVACY_CLEARED", reason="privacy_post_check_passed")
        _update_case_privacy_summary(root, approved=True, status=manifest["status"])
        return manifest
    except Exception as exc:
        if staging.exists():
            try:
                delete_tree(staging)
            except OSError:
                pass
        manifest["status"] = "BLOCKED"
        # Keep detailed errors ephemeral for the operator; the durable audit
        # record must never repeat a header, filename or removed value.
        manifest["deterministic_validation"]["errors"] = [f"PRIVACY_POST_CHECK_FAILED:{type(exc).__name__}"]
        _blocked_temporary_cleanup(
            root,
            manifest,
            delete_tree=delete_tree,
            review_file=review_file,
            known_incoming_files=incoming_files,
        )
        _atomic_json(privacy_manifest_path(root), manifest)
        _sync_lifecycle(root, "PRIVACY_BLOCKED", reason="privacy_post_check_failed")
        _update_case_privacy_summary(root, approved=False, status="BLOCKED")
        raise


@with_preservation_policy
def validate_codex_privacy_review(
    case_directory: str | Path,
    review_path: str | Path,
    *,
    delete_tree: Callable[[Path], None] = shutil.rmtree,
    supplemental: bool = False,
) -> dict[str, Any]:
    """Run the post-check after contract clearance; privacy failures remain blocking."""

    from .contract_policy import assert_contract_permission
    # Do not parse, hash or purge an unauthorized manually copied drop in this path.
    assert_contract_permission(case_directory)
    prior_manifest = None
    if supplemental:
        root = Path(privacy_requirement(case_directory)["case_root"])
        prior_manifest = _read_json(privacy_manifest_path(root))
        if (prior_manifest.get("approved_for_analysis") is not True
                or prior_manifest.get("schema_version") != MANIFEST_SCHEMA
                or prior_manifest.get("policy_version") != POLICY_VERSION
                or prior_manifest.get("status") not in {"PASS", "SANITIZED"}
                or prior_manifest.get("deterministic_validation", {}).get("passed") is not True
                or prior_manifest.get("original_deletion", {}).get("succeeded") is not True):
            raise ValueError("Additional evidence requires a previously cleared privacy batch.")
        if prior_manifest.get("business_preservation_policy") != current_preservation_policy().policy_id:
            raise ValueError("Additional evidence must retain the case's business preservation policy.")
        if prior_manifest.get("risk_policy_version") != current_preservation_policy().risk.policy_id:
            raise ValueError("Additional evidence must retain the case's privacy-risk policy.")
        approved_hashes = {entry.get("sanitized_sha256") for entry in prior_manifest["files"] if entry.get("sanitized_sha256")}
        actual_hashes = {_sha256(path) for path in _safe_files(root / "sanitized")}
        if actual_hashes != approved_hashes:
            raise ValueError("Previous sanitized evidence is missing, altered or contains unapproved files.")
    try:
        return _validate_codex_privacy_review_impl(
            case_directory, review_path, delete_tree=delete_tree, prior_manifest=prior_manifest
        )
    except Exception as exc:
        try:
            requirement = privacy_requirement(case_directory)
            if requirement.get("required") and not requirement.get("legacy"):
                root: Path = requirement["case_root"]
                if inspect_privacy_status(root)["state"] == "PURGED":
                    raise ValueError("Dossier PURGED immuable.")
                path = privacy_manifest_path(root)
                if not path.is_file():
                    files: list[dict[str, Any]] = []
                    cleanup_authorized = True
                    try:
                        incoming_files = _safe_files(root / "incoming")
                    except (OSError, ValueError):
                        incoming_files = []
                        cleanup_inventory: list[Path] | None = None
                        cleanup_authorized = False
                    else:
                        cleanup_inventory = incoming_files
                    for index, source in enumerate(incoming_files, 1):
                        try:
                            original_sha256: str | None = _sha256(source)
                        except OSError:
                            original_sha256 = None
                            cleanup_authorized = False
                        files.append({
                            "file_id": f"FILE-{index:03d}",
                            "original_sha256": original_sha256,
                            "file_type": source.suffix.casefold().lstrip(".") or "unknown",
                            "status": "BLOCKED",
                            "sanitized_sha256": None,
                        })
                    blocked_manifest = {
                        "schema_version": MANIFEST_SCHEMA,
                        "workspace_id": root.name,
                        "received_at_utc": None,
                        "validated_at_utc": _now(),
                        "policy_version": POLICY_VERSION,
                        "status": "BLOCKED",
                        "blocked_reason_codes": ["PRIVACY_REVIEW_CONTRACT_FAILED"],
                        "detected_categories": [],
                        "transformations": [],
                        "files": files,
                        "deterministic_validation": {
                            "passed": False,
                            "checks": [],
                            "errors": [f"PRIVACY_POST_CHECK_FAILED:{type(exc).__name__}"],
                        },
                        "original_deletion": {
                            "required": True,
                            "attempted": False,
                            "succeeded": False,
                            "partial": False,
                            "deleted_file_count": 0,
                            "remaining_file_count": None,
                            "deleted_at_utc": None,
                            "errors": [],
                        },
                        "approved_for_analysis": False,
                    }
                    review_candidate = Path(review_path).expanduser().resolve(strict=False)
                    _blocked_temporary_cleanup(
                        root,
                        blocked_manifest,
                        delete_tree=delete_tree,
                        review_file=review_candidate,
                        known_incoming_files=cleanup_inventory,
                        deletion_authorized=cleanup_authorized,
                    )
                    _atomic_json(path, blocked_manifest)
                else:
                    blocked_manifest = _read_json(path)
                    blocked_manifest["status"] = "BLOCKED"
                    blocked_manifest["approved_for_analysis"] = False
                    deterministic = blocked_manifest.setdefault(
                        "deterministic_validation",
                        {"passed": False, "checks": [], "errors": []},
                    )
                    deterministic["passed"] = False
                    error_code = f"PRIVACY_POST_CHECK_FAILED:{type(exc).__name__}"
                    deterministic["errors"] = list(dict.fromkeys([
                        *deterministic.get("errors", []), error_code,
                    ]))
                    previous_deletion = dict(blocked_manifest.get("original_deletion", {}))
                    try:
                        pending_files: list[Path] | None = _safe_files(root / "incoming")
                    except (OSError, ValueError):
                        pending_files = None
                    cleanup_authorized = pending_files is not None
                    if pending_files:
                        audit_files = blocked_manifest.setdefault("files", [])
                        known_hashes = {
                            item.get("original_sha256")
                            for item in audit_files if isinstance(item, dict)
                        }
                        used_file_ids = {
                            str(item.get("file_id"))
                            for item in audit_files if isinstance(item, dict)
                        }
                        next_file_number = 1
                        for source in pending_files:
                            try:
                                digest = _sha256(source)
                            except OSError:
                                cleanup_authorized = False
                                continue
                            if digest in known_hashes:
                                continue
                            while f"FILE-{next_file_number:03d}" in used_file_ids:
                                next_file_number += 1
                            file_id = f"FILE-{next_file_number:03d}"
                            audit_files.append({
                                "file_id": file_id,
                                "original_sha256": digest,
                                "file_type": source.suffix.casefold().lstrip(".") or "unknown",
                                "status": "BLOCKED",
                                "sanitized_sha256": None,
                            })
                            used_file_ids.add(file_id)
                            known_hashes.add(digest)
                            next_file_number += 1
                    _blocked_temporary_cleanup(
                        root,
                        blocked_manifest,
                        delete_tree=delete_tree,
                        review_file=Path(review_path).expanduser().resolve(strict=False),
                        known_incoming_files=pending_files,
                        deletion_authorized=cleanup_authorized,
                    )
                    if pending_files == [] and previous_deletion.get("attempted"):
                        blocked_manifest["original_deletion"] = previous_deletion
                    _atomic_json(path, blocked_manifest)
                _sync_lifecycle(root, "PRIVACY_BLOCKED", reason="privacy_validation_contract_failed")
                _update_case_privacy_summary(root, approved=False, status="BLOCKED")
        except Exception:
            # Never replace the actionable validation error with audit bookkeeping.
            pass
        raise


def configure_retention(case_directory: str | Path, policy: dict[str, Any]) -> dict[str, Any]:
    root: Path = privacy_requirement(case_directory)["case_root"]
    if policy.get("schema_version") != RETENTION_SCHEMA:
        raise ValueError("Schéma de rétention invalide.")
    if policy.get("configured") is not True or not str(policy.get("purge_after_utc", "")).strip():
        raise ValueError("La politique contractuelle doit être configurée avec purge_after_utc.")
    try:
        purge_after = datetime.fromisoformat(str(policy["purge_after_utc"]).replace("Z", "+00:00"))
    except ValueError as exc:
        raise ValueError("purge_after_utc doit être un timestamp ISO avec fuseau.") from exc
    if purge_after.tzinfo is None:
        raise ValueError("purge_after_utc doit inclure un fuseau.")
    if not isinstance(policy.get("derived_retention_authorized", False), bool):
        raise ValueError("derived_retention_authorized doit être booléen.")
    retained = policy.get("retained_paths", [])
    if not isinstance(retained, list) or any(Path(str(item)).is_absolute() or ".." in Path(str(item)).parts for item in retained):
        raise ValueError("retained_paths doit contenir uniquement des chemins relatifs sûrs.")
    allowed_prefixes = ("contracts/", "billing/", "outputs/")
    if any(not str(item).startswith(allowed_prefixes) for item in retained):
        raise ValueError("Seuls contrat, facturation et livrables explicitement autorisés peuvent être retenus.")
    retained_derived = policy.get("retained_derived_paths", [])
    if (
        not isinstance(retained_derived, list)
        or any(Path(str(item)).is_absolute() or ".." in Path(str(item)).parts for item in retained_derived)
        or any(not str(item).startswith("retained_derived/") for item in retained_derived)
    ):
        raise ValueError("retained_derived_paths doit rester dans retained_derived/.")
    if type(policy.get("mission_closed", False)) is not bool:
        raise ValueError("mission_closed doit être un booléen explicite.")
    from .contract_policy import validate_retention_authorization, contract_policy_digest, POLICY_PATH
    contract = validate_retention_authorization(root, policy)
    stored = {
        "schema_version": RETENTION_SCHEMA,
        "configured": True,
        "configured_at_utc": _now(),
        "purge_after_utc": policy["purge_after_utc"],
        "mission_closed": bool(policy.get("mission_closed", False)),
        "derived_retention_authorized": policy.get("derived_retention_authorized", False),
        "retained_paths": retained,
        "retained_derived_paths": list(retained_derived),
    }
    if contract is not None:
        stored["contract_policy_ref"] = POLICY_PATH
        stored["contract_policy_sha256"] = contract_policy_digest(contract)
    if stored["retained_derived_paths"] and not stored["derived_retention_authorized"]:
        raise ValueError("Aucun dérivé ne peut être retenu sans autorisation explicite.")
    _atomic_json(root / "privacy" / "retention_policy.json", stored)
    for name in ("workspace.json", "case_manifest.json"):
        manifest_path = root / name
        if manifest_path.is_file():
            manifest = _read_json(manifest_path)
            manifest["retention"] = {
                "configured": True,
                "policy": "privacy/retention_policy.json",
                "purge_after_utc": stored["purge_after_utc"],
                "mission_closed": stored["mission_closed"],
                "derived_retention_authorized": stored["derived_retention_authorized"],
            }
            _atomic_json(manifest_path, manifest)
            break
    return stored


def _retained_derived_allowed(root: Path, policy: dict[str, Any]) -> set[str]:
    paths = set(str(item) for item in policy.get("retained_derived_paths", []))
    if not paths:
        return set()
    review = _read_json(root / "privacy" / "derived_retention_review.json")
    required_true = ("approved_for_long_term_retention", "site_identity_removed", "unique_asset_combinations_generalized", "volumes_generalized", "timestamps_generalized")
    if any(review.get(field) is not True for field in required_true) or review.get("reidentification_risk") != "LOW":
        raise ValueError("Le dérivé ne possède pas une revue de désidentification suffisante.")
    from againward.compat.learning_retention import validate_retained_learning_projection
    from .contract_policy import assert_contract_permission, contract_policy_digest
    contract = assert_contract_permission(root, operation="retention", purpose=review.get("purpose"))
    if (review.get("status") != "RETENTION_APPROVED"
            or review.get("no_personal_data_remaining") is not True
            or review.get("no_forbidden_industrial_dimensions") is not True
            or not review.get("reviewer_role") or not review.get("reviewed_at_utc")
            or (contract is not None and review.get("contract_policy_sha256") != contract_policy_digest(contract))):
        raise ValueError("Rétention dérivée sans finalité contractuelle et revue humaine actuelles.")
    for relative in paths:
        path = root / relative
        if (
            not _within(path, root / "retained_derived")
            or not path.is_file()
            or path.suffix.casefold() not in _TABULAR_EXTENSIONS
            or _scan(path)["categories"]
        ):
            raise ValueError("Dérivé retenu absent, hors zone ou contenant des motifs sensibles.")
        row = validate_retained_learning_projection(path)
        if row["purpose"] != review.get("purpose") or review.get("artifact_sha256", {}).get(relative) != _sha256(path):
            raise ValueError("Dérivé modifié après revue ou finalité incohérente.")
        forbidden_parts = {
            "site", "machine", "equipment", "equipement", "asset", "meter", "compteur",
            "line", "ligne", "workshop", "atelier", "lot", "batch", "product", "produit",
            "reference", "timestamp", "datetime", "operator", "technician", "employee",
        }
        tables = _tables(path)
        if not tables or any(
            set(_slug(header).split("_")) & forbidden_parts
            for table in tables for header in table.headers
        ):
            raise ValueError("Le dérivé retenu expose encore une dimension industrielle réidentifiante.")
    return paths


def _retained(relative: str, protected: set[str]) -> bool:
    return any(relative == item.rstrip("/") or (item.endswith("/") and relative.startswith(item)) for item in protected)


def _purge_category(relative: str) -> str:
    first = Path(relative).parts[0] if Path(relative).parts else "root"
    if first in {"incoming", "raw"}:
        return "temporary_original"
    if first == "sanitized":
        return "sanitized_source"
    if first in {"normalized", "derived", "processed", "investigation", "evidence"}:
        return "reconstructible_derived"
    if first in {"scratch", "logs", ".cache", "cache", "tmp", "temp"}:
        return "scratch_and_cache"
    if first == "privacy":
        return "privacy_auxiliary"
    if first == "outputs":
        return "outputs"
    if first == "retained_derived":
        return "retained_derived_default_denied"
    return "unclassified_client_artifact"


def purge_client_case(
    case_directory: str | Path,
    *,
    now: datetime | None = None,
    unlink_file: Callable[[Path], None] = Path.unlink,
) -> dict[str, Any]:
    """Execute the configured end-of-mission purge and write a content-free receipt."""

    root: Path = privacy_requirement(case_directory)["case_root"]
    if inspect_privacy_status(root)["state"] == "PURGED":
        raise ValueError("Le dossier a déjà été purgé; le reçu existant reste autoritatif.")
    policy = _read_json(root / "privacy" / "retention_policy.json")
    if policy.get("schema_version") != RETENTION_SCHEMA or policy.get("configured") is not True:
        raise ValueError("Politique de rétention non configurée.")
    if policy.get("mission_closed") is not True:
        raise ValueError("La mission doit être explicitement close avant purge.")
    try:
        due = datetime.fromisoformat(str(policy["purge_after_utc"]).replace("Z", "+00:00"))
    except ValueError as exc:
        raise ValueError("purge_after_utc invalide.") from exc
    instant = now or datetime.now(timezone.utc)
    if due.tzinfo is None:
        raise ValueError("purge_after_utc doit inclure un fuseau.")
    if instant.astimezone(timezone.utc) < due.astimezone(timezone.utc):
        raise ValueError("La période contractuelle de rétention n'est pas encore échue.")
    from .contract_policy import validate_retention_authorization, contract_policy_digest
    contract = validate_retention_authorization(root, policy)
    if contract is not None and policy.get("contract_policy_sha256") != contract_policy_digest(contract):
        raise ValueError("CONTRACT_GATE: rétention à revalider après modification contractuelle.")
    retained = set(str(item) for item in policy.get("retained_paths", []))
    retained_derived = _retained_derived_allowed(root, policy) if policy.get("derived_retention_authorized") else set()
    protected = retained | retained_derived | {
        "workspace.json", "case_manifest.json",
        "privacy/privacy_manifest.json", "privacy/retention_policy.json",
    }
    results: list[dict[str, Any]] = []
    for index, path in enumerate(_safe_files(root), 1):
        relative = str(path.relative_to(root))
        if relative == "PURGE_RECEIPT.json":
            continue
        logical_id = f"PURGE-{index:05d}-{hashlib.sha256(relative.encode('utf-8')).hexdigest()[:12]}"
        category = _purge_category(relative)
        digest = _sha256(path)
        if _retained(relative, protected):
            results.append({"category": category, "logical_id": logical_id, "sha256": digest, "status": "retained_by_policy", "_relative": relative})
            continue
        try:
            unlink_file(path)
            results.append({"category": category, "logical_id": logical_id, "sha256": digest, "status": "deleted", "_relative": relative})
        except OSError as exc:
            results.append({"category": category, "logical_id": logical_id, "sha256": digest, "status": "failed", "error": type(exc).__name__, "_relative": relative})
    errors = [item for item in results if item["status"] == "failed"]
    purged_at = _now()
    privacy_manifest = {"schema_version": MANIFEST_SCHEMA, "policy_version": POLICY_VERSION}
    privacy_manifest["approved_for_analysis"] = False
    privacy_manifest["status"] = "PURGED" if not errors else "PURGE_PARTIAL_FAILURE"
    privacy_manifest["purge"] = {
        "receipt": "PURGE_RECEIPT.json",
        "completed_at_utc": purged_at,
        "succeeded": not errors,
    }
    _atomic_json(privacy_manifest_path(root), privacy_manifest)
    _update_case_privacy_summary(root, approved=False, status="PURGED" if not errors else "PURGE_PARTIAL_FAILURE")
    for name in ("workspace.json", "case_manifest.json"):
        manifest_path = root / name
        if manifest_path.is_file():
            previous = _read_json(manifest_path)
            case_manifest = {"schema_version": previous.get("schema_version"),
                "case_kind": previous.get("case_kind"), "status": "purged" if not errors else "purge_partial_failure",
                "privacy": {"required": previous.get("privacy", {}).get("required", True),
                    "policy_version": POLICY_VERSION, "approved_for_analysis": False}}
            retention = case_manifest.setdefault("retention", {})
            retention["purge_status"] = "complete" if not errors else "partial_failure"
            retention["purged_at_utc"] = purged_at
            _atomic_json(manifest_path, case_manifest)
            break
    for item in results:
        retained_path = root / item["_relative"]
        if item["status"] == "retained_by_policy" and retained_path.is_file():
            item["sha256"] = _sha256(retained_path)
    public_results = [
        {key: value for key, value in item.items() if key != "_relative"}
        for item in results
    ]
    receipt = {
        "schema_version": PURGE_SCHEMA,
        "workspace_id": root.name,
        "purged_at_utc": purged_at,
        "status": "complete" if not errors else "partial_failure",
        "derived_retention_authorized": bool(policy.get("derived_retention_authorized")),
        "items": public_results,
        "errors": [{"logical_id": item["logical_id"], "error": item["error"]} for item in errors],
        "contains_client_content": False,
    }
    _atomic_json(root / "PURGE_RECEIPT.json", receipt)
    return receipt

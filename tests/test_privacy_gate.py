from __future__ import annotations

import json
import shutil
import subprocess
from datetime import datetime, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from tests.contract_fixtures import authorize_test_case

from energy_mvp.attribution_workflow import run_minimal_attribution
from energy_mvp.client_lifecycle import (
    assert_workflow_action_allowed,
    complete_resume,
    initialize_client_lifecycle,
    publish_client_requests,
    record_canonical_answers,
    record_existing_data_exhaustion,
)
from energy_mvp.client_workspace import create_client_workspace
from energy_mvp.minimal_attribution import AnonymousElectricalComponent, EquipmentRecord
from energy_mvp.privacy import (
    POLICY_VERSION,
    RETENTION_SCHEMA,
    REVIEW_SCHEMA,
    configure_retention,
    inspect_privacy_status,
    purge_client_case,
    stage_incoming_drop,
    validate_codex_privacy_review,
)
from energy_mvp.workflow import prepare_investigation
from client_intake_pipeline import create_client_case, ingest_client_drop, record_structured_findings


def _case(tmp_path: Path, name: str = "client") -> Path:
    create_client_workspace(name, root=tmp_path / "workspaces")
    case = tmp_path / "workspaces" / name
    authorize_test_case(case)
    return case


def _file_spec(
    source: str,
    *,
    action: str = "PASS",
    sanitized: str | None = None,
    file_id: str = "FILE-001",
    transformations: list[dict] | None = None,
) -> dict:
    return {
        "file_id": file_id,
        "source": source,
        "action": action,
        "sanitized": sanitized,
        "categories": [],
        "transformations": transformations or [],
    }


def _review(
    case: Path,
    files: list[dict],
    *,
    status: str,
    categories: list[dict] | None = None,
    blocked_reasons: list[str] | None = None,
) -> Path:
    payload = {
        "schema_version": REVIEW_SCHEMA,
        "policy_version": POLICY_VERSION,
        "workspace_id": case.name,
        "received_at_utc": "2026-09-05T08:00:00+00:00",
        "status": status,
        "codex_semantic_review": {
            "completed": True,
            "first_substantive_reader_attested": True,
        },
        "detected_categories": categories or [],
        "files": files,
        "blocked_reasons": blocked_reasons or [],
    }
    path = case / "privacy" / "review.json"
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return path


def _category(name: str, action: str, count: int = 1) -> dict:
    return {"category": name, "action": action, "count": count, "file_ids": ["FILE-001"]}


def _sanitize_csv(
    case: Path,
    original: str,
    cleaned: str,
    *,
    transformations: list[dict],
    categories: list[dict],
) -> dict:
    (case / "incoming" / "data.csv").write_text(original, encoding="utf-8")
    (case / "privacy" / "candidate" / "data.csv").write_text(cleaned, encoding="utf-8")
    review = _review(
        case,
        [_file_spec(
            "incoming/data.csv",
            action="SANITIZED",
            sanitized="privacy/candidate/data.csv",
            transformations=transformations,
        )],
        status="SANITIZED",
        categories=categories,
    )
    return validate_codex_privacy_review(case, review)


def _transform(category: str, action: str, *, preserves_relations: bool = True) -> dict:
    return {
        "category": category,
        "action": action,
        "count": 1,
        "preserves_relations": preserves_relations,
        "reason_code": "ANALYSIS_MINIMIZATION",
    }


def test_clean_csv_pass_and_original_is_removed(tmp_path: Path) -> None:
    case = _case(tmp_path)
    (case / "incoming/energy.csv").write_text(
        "timestamp,machine_id,power_kw,production\n2026-01-01T00:00:00,PRESS_07,10,4\n",
        encoding="utf-8",
    )
    manifest = validate_codex_privacy_review(
        case,
        _review(case, [_file_spec("incoming/energy.csv")], status="PASS"),
    )
    assert manifest["status"] == "PASS"
    assert manifest["approved_for_analysis"] is True
    assert manifest["original_deletion"]["succeeded"] is True
    assert (case / "sanitized/energy.csv").is_file()
    assert list((case / "incoming").iterdir()) == []
    assert not (case / "privacy/review.json").exists()


def test_operator_name_is_pseudonymized_and_machine_join_preserved(tmp_path: Path) -> None:
    case = _case(tmp_path)
    manifest = _sanitize_csv(
        case,
        "timestamp,machine_id,power_kw,operator,shift\nT1,PRESS_07,10,Jean Dupont,nuit\nT2,PRESS_07,11,Jean Dupont,nuit\n",
        "timestamp,machine_id,power_kw,operator,shift\nT1,PRESS_07,10,OPERATOR_001,nuit\nT2,PRESS_07,11,OPERATOR_001,nuit\n",
        transformations=[_transform("PERSON_NAME", "PSEUDONYMIZED")],
        categories=[_category("PERSON_NAME", "PSEUDONYMIZED", 2)],
    )
    text = (case / "sanitized/data.csv").read_text(encoding="utf-8")
    assert "OPERATOR_001" in text and "PRESS_07" in text
    assert "Jean Dupont" not in text
    assert manifest["deterministic_validation"]["passed"] is True


@pytest.mark.parametrize(
    ("column", "value", "category"),
    [
        ("email", "jean.dupont@example.com", "EMAIL"),
        ("telephone", "+33 6 12 34 56 78", "PHONE"),
    ],
)
def test_email_and_phone_are_removed(tmp_path: Path, column: str, value: str, category: str) -> None:
    case = _case(tmp_path)
    _sanitize_csv(
        case,
        f"timestamp,power_kw,{column}\nT1,10,{value}\n",
        f"timestamp,power_kw\nT1,10\n",
        transformations=[_transform(category, "REMOVED")],
        categories=[_category(category, "REMOVED")],
    )
    assert value not in (case / "sanitized/data.csv").read_text(encoding="utf-8")


def test_employee_number_is_stably_pseudonymized(tmp_path: Path) -> None:
    case = _case(tmp_path)
    _sanitize_csv(
        case,
        "timestamp,machine_id,matricule,power_kw\nT1,M1,84721,10\nT2,M2,84721,11\n",
        "timestamp,machine_id,matricule,power_kw\nT1,M1,OPERATOR_001,10\nT2,M2,OPERATOR_001,11\n",
        transformations=[_transform("EMPLOYEE_ID", "PSEUDONYMIZED")],
        categories=[_category("EMPLOYEE_ID", "PSEUDONYMIZED", 2)],
    )


def test_secret_must_be_removed_and_residual_secret_blocks(tmp_path: Path) -> None:
    case = _case(tmp_path, "clean_secret")
    _sanitize_csv(
        case,
        "timestamp,power_kw,api_key\nT1,10,sk-live-1234567890abcdef\n",
        "timestamp,power_kw\nT1,10\n",
        transformations=[_transform("AUTHENTICATION_SECRET", "REMOVED")],
        categories=[_category("AUTHENTICATION_SECRET", "REMOVED")],
    )
    blocked = _case(tmp_path, "blocked_secret")
    (blocked / "incoming/data.csv").write_text(
        "timestamp,power_kw,notes\nT1,10,sk-live-1234567890abcdef\n", encoding="utf-8"
    )
    (blocked / "privacy/candidate/data.csv").write_text(
        "timestamp,power_kw,notes\nT1,10,sk-live-1234567890abcdef\n", encoding="utf-8"
    )
    review = _review(
        blocked,
        [_file_spec("incoming/data.csv", action="SANITIZED", sanitized="privacy/candidate/data.csv", transformations=[_transform("AUTHENTICATION_SECRET", "REMOVED")])],
        status="SANITIZED",
        categories=[_category("AUTHENTICATION_SECRET", "REMOVED")],
    )
    with pytest.raises(ValueError, match="post-check"):
        validate_codex_privacy_review(blocked, review)
    assert inspect_privacy_status(blocked)["state"] == "PRIVACY_BLOCKED"


def test_medical_or_hr_content_is_fail_closed(tmp_path: Path) -> None:
    case = _case(tmp_path)
    (case / "incoming/rh.txt").write_text(
        "Dossier médical et sanction disciplinaire concernant un salarié.", encoding="utf-8"
    )
    manifest = validate_codex_privacy_review(
        case,
        _review(
            case,
            [_file_spec("incoming/rh.txt", action="BLOCKED")],
            status="BLOCKED",
            categories=[_category("MEDICAL_DATA", "BLOCKED")],
            blocked_reasons=["SPECIAL_CATEGORY_DATA"],
        ),
    )
    assert manifest["approved_for_analysis"] is False
    assert manifest["original_deletion"]["succeeded"] is True
    assert manifest["original_deletion"]["deleted_file_count"] == 1
    assert not (case / "incoming").exists()
    assert not (case / "privacy/review.json").exists()
    assert initialize_client_lifecycle(case)["client_lifecycle"]["state"] == "PRIVACY_BLOCKED"


def test_blocked_deletion_failure_is_reported_honestly(tmp_path: Path) -> None:
    case = _case(tmp_path)
    raw = case / "incoming/rh.txt"
    raw.write_text("Dossier médical nécessitant un arrêt immédiat.", encoding="utf-8")

    def fail_incoming(path: Path) -> None:
        if path.name == "incoming":
            raise PermissionError("simulated")
        shutil.rmtree(path)

    manifest = validate_codex_privacy_review(
        case,
        _review(
            case,
            [_file_spec("incoming/rh.txt", action="BLOCKED")],
            status="BLOCKED",
            categories=[_category("MEDICAL_DATA", "BLOCKED")],
            blocked_reasons=["SPECIAL_CATEGORY_DATA"],
        ),
        delete_tree=fail_incoming,
    )

    assert manifest["status"] == "BLOCKED"
    assert manifest["approved_for_analysis"] is False
    assert manifest["original_deletion"]["attempted"] is True
    assert manifest["original_deletion"]["succeeded"] is False
    assert manifest["original_deletion"]["remaining_file_count"] == 1
    assert manifest["original_deletion"]["deleted_at_utc"] is None
    assert raw.is_file()
    durable = json.loads((case / "privacy/privacy_manifest.json").read_text())
    assert durable["original_deletion"]["succeeded"] is False
    assert "Dossier médical" not in json.dumps(durable, ensure_ascii=False)


def test_blocked_partial_deletion_reports_remaining_count(tmp_path: Path) -> None:
    case = _case(tmp_path)
    for name in ("a.txt", "b.txt"):
        (case / "incoming" / name).write_text("Données RH sensibles.", encoding="utf-8")

    def delete_one_then_fail(path: Path) -> None:
        if path.name == "incoming":
            (path / "a.txt").unlink()
            raise PermissionError("simulated partial deletion")
        shutil.rmtree(path)

    manifest = validate_codex_privacy_review(
        case,
        _review(
            case,
            [
                _file_spec("incoming/a.txt", action="BLOCKED", file_id="FILE-001"),
                _file_spec("incoming/b.txt", action="BLOCKED", file_id="FILE-002"),
            ],
            status="BLOCKED",
            categories=[{
                "category": "HR_SENSITIVE", "action": "BLOCKED", "count": 2,
                "file_ids": ["FILE-001", "FILE-002"],
            }],
            blocked_reasons=["SPECIAL_CATEGORY_DATA"],
        ),
        delete_tree=delete_one_then_fail,
    )

    deletion = manifest["original_deletion"]
    assert deletion["succeeded"] is False
    assert deletion["partial"] is True
    assert deletion["deleted_file_count"] == 1
    assert deletion["remaining_file_count"] == 1
    assert not (case / "incoming/a.txt").exists()
    assert (case / "incoming/b.txt").exists()


def test_blocked_case_can_stage_a_fresh_minimal_drop_after_raw_purge(tmp_path: Path) -> None:
    case = _case(tmp_path)
    (case / "incoming/rh.txt").write_text("Dossier médical.", encoding="utf-8")
    validate_codex_privacy_review(
        case,
        _review(
            case,
            [_file_spec("incoming/rh.txt", action="BLOCKED")],
            status="BLOCKED",
            categories=[_category("MEDICAL_DATA", "BLOCKED")],
            blocked_reasons=["SPECIAL_CATEGORY_DATA"],
        ),
    )
    fresh = tmp_path / "fresh"
    fresh.mkdir()
    (fresh / "energy.csv").write_text(
        "timestamp,power_kw\nT1,10\n", encoding="utf-8"
    )

    receipt = stage_incoming_drop(fresh, case)

    assert receipt["file_count"] == 1
    assert (case / "incoming/energy.csv").is_file()
    assert inspect_privacy_status(case)["state"] == "PRIVACY_BLOCKED"
    manifest = validate_codex_privacy_review(
        case, _review(case, [_file_spec("incoming/energy.csv")], status="PASS")
    )
    assert manifest["approved_for_analysis"] is True
    assert inspect_privacy_status(case)["state"] == "PRIVACY_CLEARED"


def test_industrial_values_timestamps_and_production_are_exact(tmp_path: Path) -> None:
    case = _case(tmp_path)
    original = (
        "timestamp,machine_name,meter_id,energy_kwh,power_kw,production,product_type,operator\n"
        "2026-01-01T00:00:00,Jean,MTR-1,2.500,10.00,4,Jean,Paul Martin\n"
    )
    cleaned = (
        "timestamp,machine_name,meter_id,energy_kwh,power_kw,production,product_type,operator\n"
        "2026-01-01T00:00:00,Jean,MTR-1,2.500,10.00,4,Jean,OPERATOR_001\n"
    )
    _sanitize_csv(
        case,
        original,
        cleaned,
        transformations=[_transform("PERSON_NAME", "PSEUDONYMIZED")],
        categories=[_category("PERSON_NAME", "PSEUDONYMIZED")],
    )
    output = (case / "sanitized/data.csv").read_text(encoding="utf-8")
    for value in ("2026-01-01T00:00:00", "Jean", "MTR-1", "2.500", "10.00", ",4,"):
        assert value in output


def test_manifest_never_contains_removed_values(tmp_path: Path) -> None:
    case = _case(tmp_path)
    manifest = _sanitize_csv(
        case,
        "timestamp,power_kw,operator,email\nT1,10,Jean Dupont,jean@example.com\n",
        "timestamp,power_kw,operator\nT1,10,OPERATOR_001\n",
        transformations=[_transform("PERSON_NAME", "PSEUDONYMIZED"), _transform("EMAIL", "REMOVED")],
        categories=[_category("PERSON_NAME", "PSEUDONYMIZED"), _category("EMAIL", "REMOVED")],
    )
    serialized = json.dumps(manifest, ensure_ascii=False)
    assert "Jean Dupont" not in serialized
    assert "jean@example.com" not in serialized


def test_analysis_is_impossible_before_clearance_and_possible_after(tmp_path: Path) -> None:
    case = _case(tmp_path)
    source = case / "incoming/energy.csv"
    source.write_text("timestamp,energy_kwh\n2026-01-01,10\n2026-01-02,11\n", encoding="utf-8")
    with pytest.raises(ValueError, match="PRIVACY GATE"):
        prepare_investigation(source, case / "processed")
    validate_codex_privacy_review(case, _review(case, [_file_spec("incoming/energy.csv")], status="PASS"))
    state = prepare_investigation(case / "sanitized/energy.csv", case / "processed")
    assert state["client_lifecycle"]["state"] == "ANALYZING"


def test_real_case_cannot_write_derived_outputs_outside_its_workspace(tmp_path: Path) -> None:
    case = _case(tmp_path)
    (case / "incoming/energy.csv").write_text(
        "timestamp,energy_kwh\n2026-01-01,10\n2026-01-02,11\n", encoding="utf-8"
    )
    validate_codex_privacy_review(case, _review(case, [_file_spec("incoming/energy.csv")], status="PASS"))
    with pytest.raises(ValueError, match="même workspace"):
        prepare_investigation(case / "sanitized/energy.csv", tmp_path / "escaped")


def test_sanitized_source_tampering_invalidates_clearance_for_analysis(tmp_path: Path) -> None:
    case = _case(tmp_path)
    (case / "incoming/energy.csv").write_text(
        "timestamp,energy_kwh\n2026-01-01,10\n", encoding="utf-8"
    )
    validate_codex_privacy_review(case, _review(case, [_file_spec("incoming/energy.csv")], status="PASS"))
    (case / "sanitized/energy.csv").write_text(
        "timestamp,energy_kwh\n2026-01-01,99\n", encoding="utf-8"
    )
    with pytest.raises(ValueError, match="hash"):
        prepare_investigation(case / "sanitized/energy.csv", case / "processed")


def test_any_review_contract_failure_persists_privacy_blocked(tmp_path: Path) -> None:
    case = _case(tmp_path)
    (case / "incoming/energy.csv").write_text(
        "timestamp,power_kw\nT1,10\n", encoding="utf-8"
    )
    review = json.loads(_review(case, [_file_spec("incoming/energy.csv")], status="PASS").read_text())
    review["schema_version"] = "wrong"
    path = case / "privacy/review.json"
    path.write_text(json.dumps(review), encoding="utf-8")
    with pytest.raises(ValueError, match="Schéma"):
        validate_codex_privacy_review(case, path)
    status = inspect_privacy_status(case)
    assert status["state"] == "PRIVACY_BLOCKED"
    assert status["approved_for_analysis"] is False
    assert not (case / "incoming").exists()


def test_pseudonym_correspondence_table_outside_candidate_is_refused(tmp_path: Path) -> None:
    case = _case(tmp_path)
    (case / "incoming/energy.csv").write_text(
        "timestamp,power_kw\nT1,10\n", encoding="utf-8"
    )
    (case / "privacy/pseudonym_map.json").write_text(
        '{"OPERATOR_001": "Jean Dupont"}', encoding="utf-8"
    )
    with pytest.raises(ValueError, match="auxiliaire inattendu"):
        validate_codex_privacy_review(
            case, _review(case, [_file_spec("incoming/energy.csv")], status="PASS")
        )
    assert inspect_privacy_status(case)["state"] == "PRIVACY_BLOCKED"


def test_goal_a_intake_reads_sanitized_not_incoming(tmp_path: Path) -> None:
    create_client_case("goal_a", root=tmp_path / "cases")
    case = tmp_path / "cases/goal_a"
    authorize_test_case(case)
    (case / "incoming/energy.csv").write_text(
        "timestamp,energy_kwh\n2026-01-01T00:00:00,2\n2026-01-01T00:30:00,3\n",
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="PRIVACY GATE"):
        ingest_client_drop(case / "incoming", case)
    validate_codex_privacy_review(
        case, _review(case, [_file_spec("incoming/energy.csv")], status="PASS")
    )
    canonical = ingest_client_drop(case / "sanitized", case)
    assert canonical["available_datasets"]
    provenance = json.loads((case / "evidence/intake_inventory.json").read_text())
    assert provenance["artifacts"][0]["raw_relative_path"].startswith("sanitized/")
    assert not (case / "raw").exists()


def test_staging_is_content_blind_and_records_only_logical_paths(tmp_path: Path) -> None:
    case = _case(tmp_path)
    drop = tmp_path / "drop"
    drop.mkdir()
    (drop / "client.csv").write_text("Jean Dupont,secret-value\n", encoding="utf-8")
    receipt = stage_incoming_drop(drop, case)
    assert receipt["content_semantically_inspected"] is False
    assert "sha256" not in receipt
    assert "Jean Dupont" not in json.dumps(receipt)
    assert (case / "incoming/client.csv").read_text(encoding="utf-8") == "Jean Dupont,secret-value\n"


def test_legacy_workspace_requires_explicit_privacy_migration(tmp_path: Path) -> None:
    case = tmp_path / "legacy"
    (case / "input").mkdir(parents=True)
    (case / "processed").mkdir()
    (case / "workspace.json").write_text(
        json.dumps({"schema_version": 2, "workspace_id": "legacy"}), encoding="utf-8"
    )
    source = case / "input/data.csv"
    source.write_text("timestamp,energy_kwh\n2026-01-01,10\n", encoding="utf-8")
    with pytest.raises(ValueError, match="migration privacy"):
        prepare_investigation(source, case / "processed")


def test_answer_resume_cycle_remains_functional_after_privacy_clearance(tmp_path: Path) -> None:
    case = _case(tmp_path)
    (case / "incoming/energy.csv").write_text(
        "timestamp,energy_kwh\n2026-01-01,10\n2026-01-02,11\n", encoding="utf-8"
    )
    validate_codex_privacy_review(case, _review(case, [_file_spec("incoming/energy.csv")], status="PASS"))
    prepare_investigation(case / "sanitized/energy.csv", case / "processed")
    record_existing_data_exhaustion(case, analysis_inventory_ref="prepared_analysis.json", reviewed_sources=["evidence_card.json"])
    candidate = {
        "request_id": "Q1", "request_type": "MICRO_QUESTION",
        "client_question": "Le site était-il fermé pendant ce palier ?",
        "internal_reason": "Départage activité utile et charge inexpliquée.",
        "target_role": "responsable de site", "related_hypothesis_ids": ["H1"],
        "related_finding_ids": ["F1"], "related_component_ids": [],
        "hypotheses_distinguished": ["activité utile", "charge inexpliquée"],
        "plausible_answers": [
            {"answer_id": "OPEN", "label": "Ouvert", "decision_effects": ["écarte la piste"]},
            {"answer_id": "CLOSED", "label": "Fermé", "decision_effects": ["renforce la piste"]},
        ],
        "decision_impact_dimensions": ["evidence_level", "false_conclusion_risk"],
        "effort": 1, "availability": .9, "reliability": .7, "source_cost": 0,
        "expected_source_type": "CLIENT_DECLARATION", "importance": "BLOCKING",
    }
    assert publish_client_requests(case, [candidate])["must_stop"] is True
    answer = {
        "answer_id": "A1", "request_id": "Q1", "answer": "Le site était fermé.",
        "provided_by_role": "responsable de site", "source_or_evidence": "réponse datée",
        "source_type": "CLIENT_DECLARATION", "provided_at_utc": "2026-09-05T10:00:00+02:00",
        "reliability": .65,
    }
    assert record_canonical_answers(case, [answer])["state"] == "RESUMING"
    dimensions = {
        key: "réévalué" for key in (
            "evidence_level", "asset_attribution", "alternatives", "confidence",
            "economic_materiality", "investigation_priority", "field_action",
            "false_conclusion_risk",
        )
    }
    state = complete_resume(
        case,
        recalculation_refs=["scratch/recalculation.json"],
        adversarial_review_ref="review.json",
        before_after=[{"hypothesis_id": "H1", "before": "unknown", "after": "revised", "decision_dimensions": dimensions}],
    )
    assert state["client_lifecycle"]["state"] == "ANALYZING"


def test_post_check_blocks_industrial_mutation(tmp_path: Path) -> None:
    case = _case(tmp_path)
    (case / "incoming/data.csv").write_text(
        "timestamp,machine_id,power_kw,operator\nT1,M1,10,Jean Dupont\n", encoding="utf-8"
    )
    (case / "privacy/candidate/data.csv").write_text(
        "timestamp,machine_id,power_kw,operator\nT1,M1,99,OPERATOR_001\n", encoding="utf-8"
    )
    review = _review(
        case,
        [_file_spec("incoming/data.csv", action="SANITIZED", sanitized="privacy/candidate/data.csv", transformations=[_transform("PERSON_NAME", "PSEUDONYMIZED")])],
        status="SANITIZED",
        categories=[_category("PERSON_NAME", "PSEUDONYMIZED")],
    )
    with pytest.raises(ValueError, match="modifiée"):
        validate_codex_privacy_review(case, review)
    assert inspect_privacy_status(case)["state"] == "PRIVACY_BLOCKED"


def test_block_manifest_error_never_repeats_sensitive_header_or_value(tmp_path: Path) -> None:
    case = _case(tmp_path)
    (case / "incoming/data.csv").write_text(
        "timestamp,power_kw,operator\nT1,10,Jean Dupont\n", encoding="utf-8"
    )
    (case / "privacy/candidate/data.csv").write_text(
        "timestamp,power_kw,operator\nT1,99,OPERATOR_001\n", encoding="utf-8"
    )
    review = _review(
        case,
        [_file_spec("incoming/data.csv", action="SANITIZED", sanitized="privacy/candidate/data.csv", transformations=[_transform("PERSON_NAME", "PSEUDONYMIZED")])],
        status="SANITIZED",
        categories=[_category("PERSON_NAME", "PSEUDONYMIZED")],
    )
    with pytest.raises(ValueError):
        validate_codex_privacy_review(case, review)
    manifest_text = (case / "privacy/privacy_manifest.json").read_text(encoding="utf-8")
    assert "Jean Dupont" not in manifest_text
    assert "operator" not in manifest_text


def test_failed_original_deletion_never_reports_success(tmp_path: Path) -> None:
    case = _case(tmp_path)
    (case / "incoming/energy.csv").write_text("timestamp,power_kw\nT1,10\n", encoding="utf-8")

    def fail_incoming(path: Path) -> None:
        if path.name == "incoming":
            raise PermissionError("simulated")
        shutil.rmtree(path)

    manifest = validate_codex_privacy_review(
        case,
        _review(case, [_file_spec("incoming/energy.csv")], status="PASS"),
        delete_tree=fail_incoming,
    )
    assert manifest["status"] == "BLOCKED"
    assert manifest["approved_for_analysis"] is False
    assert manifest["original_deletion"]["succeeded"] is False


def test_technician_names_are_pseudonymized_but_machine_named_jean_is_not(tmp_path: Path) -> None:
    case = _case(tmp_path)
    _sanitize_csv(
        case,
        "timestamp,machine_name,technician,power_kw\nT1,Jean,Marie Durand,10\n",
        "timestamp,machine_name,technician,power_kw\nT1,Jean,TECHNICIAN_001,10\n",
        transformations=[_transform("PERSON_NAME", "PSEUDONYMIZED")],
        categories=[_category("PERSON_NAME", "PSEUDONYMIZED")],
    )
    assert ",Jean,TECHNICIAN_001," in (case / "sanitized/data.csv").read_text(encoding="utf-8")


def test_industrial_identifier_that_looks_like_phone_is_preserved(tmp_path: Path) -> None:
    case = _case(tmp_path)
    value = "06 12 34 56 78"
    (case / "incoming/data.csv").write_text(
        f"timestamp,machine_id,power_kw\nT1,{value},10\n", encoding="utf-8"
    )
    manifest = validate_codex_privacy_review(
        case, _review(case, [_file_spec("incoming/data.csv")], status="PASS")
    )
    assert manifest["status"] == "PASS"
    assert value in (case / "sanitized/data.csv").read_text(encoding="utf-8")


def test_operator_pseudonym_keeps_cross_machine_relationship(tmp_path: Path) -> None:
    case = _case(tmp_path)
    _sanitize_csv(
        case,
        "timestamp,machine_id,operator,power_kw\nT1,M1,Alice Martin,10\nT2,M2,Alice Martin,11\n",
        "timestamp,machine_id,operator,power_kw\nT1,M1,OPERATOR_001,10\nT2,M2,OPERATOR_001,11\n",
        transformations=[_transform("PERSON_NAME", "PSEUDONYMIZED")],
        categories=[_category("PERSON_NAME", "PSEUDONYMIZED", 2)],
    )
    assert (case / "sanitized/data.csv").read_text().count("OPERATOR_001") == 2


def test_same_identity_in_two_files_requires_same_pseudonym(tmp_path: Path) -> None:
    case = _case(tmp_path)
    for name, machine in (("a.csv", "M1"), ("b.csv", "M2")):
        (case / "incoming" / name).write_text(
            f"timestamp,machine_id,operator,power_kw\nT1,{machine},Alice Martin,10\n",
            encoding="utf-8",
        )
        (case / "privacy/candidate" / name).write_text(
            f"timestamp,machine_id,operator,power_kw\nT1,{machine},OPERATOR_001,10\n",
            encoding="utf-8",
        )
    specs = [
        _file_spec(
            f"incoming/{name}",
            action="SANITIZED",
            sanitized=f"privacy/candidate/{name}",
            file_id=f"FILE-00{index}",
            transformations=[_transform("PERSON_NAME", "PSEUDONYMIZED")],
        )
        for index, name in enumerate(("a.csv", "b.csv"), 1)
    ]
    manifest = validate_codex_privacy_review(
        case,
        _review(
            case,
            specs,
            status="SANITIZED",
            categories=[{
                "category": "PERSON_NAME",
                "action": "PSEUDONYMIZED",
                "count": 2,
                "file_ids": ["FILE-001", "FILE-002"],
            }],
        ),
    )
    assert manifest["approved_for_analysis"] is True
    assert "M1,OPERATOR_001" in (case / "sanitized/a.csv").read_text()
    assert "M2,OPERATOR_001" in (case / "sanitized/b.csv").read_text()


def test_same_identity_with_different_cross_file_pseudonyms_is_blocked(tmp_path: Path) -> None:
    case = _case(tmp_path)
    specs = []
    for index, (name, pseudonym) in enumerate(
        (("a.csv", "OPERATOR_001"), ("b.csv", "OPERATOR_002")), 1
    ):
        (case / "incoming" / name).write_text(
            "timestamp,machine_id,operator,power_kw\nT1,M1,Alice Martin,10\n",
            encoding="utf-8",
        )
        (case / "privacy/candidate" / name).write_text(
            f"timestamp,machine_id,operator,power_kw\nT1,M1,{pseudonym},10\n",
            encoding="utf-8",
        )
        specs.append(_file_spec(
            f"incoming/{name}", action="SANITIZED",
            sanitized=f"privacy/candidate/{name}", file_id=f"FILE-00{index}",
            transformations=[_transform("PERSON_NAME", "PSEUDONYMIZED")],
        ))
    with pytest.raises(ValueError, match="inter-fichiers instable"):
        validate_codex_privacy_review(
            case,
            _review(case, specs, status="SANITIZED", categories=[{
                "category": "PERSON_NAME", "action": "PSEUDONYMIZED",
                "count": 2, "file_ids": ["FILE-001", "FILE-002"],
            }]),
        )
    assert inspect_privacy_status(case)["state"] == "PRIVACY_BLOCKED"
    assert not (case / "incoming").exists()


def test_two_identities_cannot_share_a_cross_file_pseudonym(tmp_path: Path) -> None:
    case = _case(tmp_path)
    specs = []
    for index, (name, identity) in enumerate(
        (("a.csv", "Alice Martin"), ("b.csv", "Bob Durand")), 1
    ):
        (case / "incoming" / name).write_text(
            f"timestamp,machine_id,operator,power_kw\nT1,M{index},{identity},10\n",
            encoding="utf-8",
        )
        (case / "privacy/candidate" / name).write_text(
            f"timestamp,machine_id,operator,power_kw\nT1,M{index},OPERATOR_001,10\n",
            encoding="utf-8",
        )
        specs.append(_file_spec(
            f"incoming/{name}", action="SANITIZED",
            sanitized=f"privacy/candidate/{name}", file_id=f"FILE-00{index}",
            transformations=[_transform("PERSON_NAME", "PSEUDONYMIZED")],
        ))
    with pytest.raises(ValueError, match="Collision de pseudonymes inter-fichiers"):
        validate_codex_privacy_review(
            case,
            _review(case, specs, status="SANITIZED", categories=[{
                "category": "PERSON_NAME", "action": "PSEUDONYMIZED",
                "count": 2, "file_ids": ["FILE-001", "FILE-002"],
            }]),
        )
    assert inspect_privacy_status(case)["state"] == "PRIVACY_BLOCKED"


@pytest.mark.parametrize("managed_root", ["workspaces", "client_cases"])
def test_unregistered_direct_path_below_managed_workspaces_cannot_bypass_gate(
    tmp_path: Path, managed_root: str
) -> None:
    case = tmp_path / managed_root / "legacy_real"
    source = case / "input" / "energy.csv"
    source.parent.mkdir(parents=True)
    source.write_text(
        "timestamp,energy_kwh\n2026-01-01,10\n2026-01-02,11\n",
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="migration privacy"):
        prepare_investigation(source, case / "processed")
    with pytest.raises(ValueError, match="migration privacy"):
        initialize_client_lifecycle(case)
    assert not (case / "investigation_state.json").exists()


def test_approved_flag_cannot_bypass_incomplete_manifest_contract(tmp_path: Path) -> None:
    case = _case(tmp_path)
    (case / "incoming/energy.csv").write_text(
        "timestamp,energy_kwh\n2026-01-01,10\n", encoding="utf-8"
    )
    validate_codex_privacy_review(
        case, _review(case, [_file_spec("incoming/energy.csv")], status="PASS")
    )
    path = case / "privacy/privacy_manifest.json"
    manifest = json.loads(path.read_text())
    manifest["original_deletion"]["succeeded"] = False
    path.write_text(json.dumps(manifest), encoding="utf-8")

    with pytest.raises(ValueError, match="PRIVACY_BLOCKED"):
        prepare_investigation(case / "sanitized/energy.csv", case / "processed")


def test_explicit_synthetic_workspace_exemption_still_works(tmp_path: Path) -> None:
    root = tmp_path / "workspaces"
    create_client_workspace("synthetic_benchmark", root=root, synthetic=True)
    case = root / "synthetic_benchmark"
    source = case / "incoming" / "energy.csv"
    source.write_text(
        "timestamp,energy_kwh\n2026-01-01,10\n2026-01-02,11\n",
        encoding="utf-8",
    )
    state = prepare_investigation(source, case / "processed")
    assert state["client_lifecycle"]["state"] == "ANALYZING"


def test_pseudonymization_cannot_silently_drop_a_relational_value(tmp_path: Path) -> None:
    case = _case(tmp_path)
    (case / "incoming/data.csv").write_text(
        "timestamp,machine_id,operator,power_kw\nT1,M1,Alice Martin,10\nT2,M2,Alice Martin,11\n",
        encoding="utf-8",
    )
    (case / "privacy/candidate/data.csv").write_text(
        "timestamp,machine_id,operator,power_kw\nT1,M1,OPERATOR_001,10\nT2,M2,,11\n",
        encoding="utf-8",
    )
    review = _review(
        case,
        [_file_spec("incoming/data.csv", action="SANITIZED", sanitized="privacy/candidate/data.csv", transformations=[_transform("PERSON_NAME", "PSEUDONYMIZED")])],
        status="SANITIZED",
        categories=[_category("PERSON_NAME", "PSEUDONYMIZED", 2)],
    )
    with pytest.raises(ValueError, match="Relation industrielle perdue"):
        validate_codex_privacy_review(case, review)


def test_free_text_email_is_removed_without_turning_privacy_into_data_cleaning(tmp_path: Path) -> None:
    case = _case(tmp_path)
    _sanitize_csv(
        case,
        "timestamp,power_kw,notes\nT1,10,Appeler jean@example.com après cycle C7\n",
        "timestamp,power_kw,notes\nT1,10,Contact supprimé après cycle C7\n",
        transformations=[_transform("EMAIL", "REMOVED")],
        categories=[_category("EMAIL", "REMOVED")],
    )
    assert "cycle C7" in (case / "sanitized/data.csv").read_text(encoding="utf-8")


def test_multisheet_xlsx_preserves_sheets_and_industrial_cells(tmp_path: Path) -> None:
    case = _case(tmp_path)
    original = Workbook()
    first = original.active; first.title = "Energy"
    first.append(["timestamp", "machine_id", "power_kw", "operator"])
    first.append(["T1", "M1", 10.25, "Jean Dupont"])
    second = original.create_sheet("Production")
    second.append(["timestamp", "production", "lot"]); second.append(["T1", 4, "L7"])
    original.save(case / "incoming/book.xlsx")
    cleaned = Workbook()
    first = cleaned.active; first.title = "Energy"
    first.append(["timestamp", "machine_id", "power_kw", "operator"])
    first.append(["T1", "M1", 10.25, "OPERATOR_001"])
    second = cleaned.create_sheet("Production")
    second.append(["timestamp", "production", "lot"]); second.append(["T1", 4, "L7"])
    cleaned.save(case / "privacy/candidate/book.xlsx")
    validate_codex_privacy_review(
        case,
        _review(
            case,
            [_file_spec("incoming/book.xlsx", action="SANITIZED", sanitized="privacy/candidate/book.xlsx", transformations=[_transform("PERSON_NAME", "PSEUDONYMIZED")])],
            status="SANITIZED",
            categories=[_category("PERSON_NAME", "PSEUDONYMIZED")],
        ),
    )
    book = load_workbook(case / "sanitized/book.xlsx", read_only=True, data_only=True)
    try:
        assert book.sheetnames == ["Energy", "Production"]
        assert book["Energy"]["C2"].value == 10.25
        assert book["Production"]["B2"].value == 4
    finally:
        book.close()


def test_text_file_and_duplicate_files_are_supported(tmp_path: Path) -> None:
    case = _case(tmp_path)
    text = "Courriel jean@example.com; machine PRESS_07, cycle nuit."
    for name in ("a.txt", "b.txt"):
        (case / "incoming" / name).write_text(text, encoding="utf-8")
        (case / "privacy/candidate" / name).write_text("Contact supprimé; machine PRESS_07, cycle nuit.", encoding="utf-8")
    specs = [
        _file_spec(f"incoming/{name}", action="SANITIZED", sanitized=f"privacy/candidate/{name}", file_id=f"FILE-00{i}", transformations=[_transform("EMAIL", "REMOVED")])
        for i, name in enumerate(("a.txt", "b.txt"), 1)
    ]
    manifest = validate_codex_privacy_review(
        case,
        _review(case, specs, status="SANITIZED", categories=[{"category": "EMAIL", "action": "REMOVED", "count": 2, "file_ids": ["FILE-001", "FILE-002"]}]),
    )
    assert manifest["files"][1]["duplicate_of_file_id"] == "FILE-001"


def test_text_sanitation_cannot_drop_machine_or_cycle_markers(tmp_path: Path) -> None:
    case = _case(tmp_path)
    (case / "incoming/note.txt").write_text(
        "Email jean@example.com; machine PRESS_07; cycle C7.", encoding="utf-8"
    )
    (case / "privacy/candidate/note.txt").write_text(
        "Contact supprimé.", encoding="utf-8"
    )
    review = _review(
        case,
        [_file_spec("incoming/note.txt", action="SANITIZED", sanitized="privacy/candidate/note.txt", transformations=[_transform("EMAIL", "REMOVED")])],
        status="SANITIZED",
        categories=[_category("EMAIL", "REMOVED")],
    )
    with pytest.raises(ValueError, match="marqueur industriel textuel"):
        validate_codex_privacy_review(case, review)


def test_large_csv_post_check_is_linear_enough_for_termux(tmp_path: Path) -> None:
    case = _case(tmp_path)
    rows = ["timestamp,machine_id,power_kw,production"]
    rows.extend(f"2026-01-01T00:{i % 60:02d}:00,M{i % 7},{10 + i % 3},{i % 5}" for i in range(50_000))
    (case / "incoming/large.csv").write_text("\n".join(rows) + "\n", encoding="utf-8")
    manifest = validate_codex_privacy_review(
        case, _review(case, [_file_spec("incoming/large.csv")], status="PASS")
    )
    assert manifest["approved_for_analysis"] is True


def test_interrupted_or_incomplete_cleaning_is_blocked(tmp_path: Path) -> None:
    case = _case(tmp_path)
    (case / "incoming/data.csv").write_text("timestamp,power_kw,operator\nT1,10,Jean Dupont\n", encoding="utf-8")
    (case / "privacy/candidate/data.csv").write_text("timestamp,power_kw,operator\n", encoding="utf-8")
    review = _review(
        case,
        [_file_spec("incoming/data.csv", action="SANITIZED", sanitized="privacy/candidate/data.csv", transformations=[_transform("PERSON_NAME", "PSEUDONYMIZED")])],
        status="SANITIZED",
        categories=[_category("PERSON_NAME", "PSEUDONYMIZED")],
    )
    with pytest.raises(ValueError, match="nombre de lignes"):
        validate_codex_privacy_review(case, review)
    assert inspect_privacy_status(case)["state"] == "PRIVACY_BLOCKED"


def test_final_purge_removes_client_data_but_retains_explicit_contract_and_report(tmp_path: Path) -> None:
    case = _case(tmp_path)
    (case / "incoming/energy.csv").write_text("timestamp,power_kw\nT1,10\n", encoding="utf-8")
    validate_codex_privacy_review(case, _review(case, [_file_spec("incoming/energy.csv")], status="PASS"))
    (case / "processed/cache.json").write_text('{"client": "value"}', encoding="utf-8")
    (case / "scratch/temp.txt").write_text("client temporary", encoding="utf-8")
    (case / "outputs/final.pdf").write_bytes(b"synthetic report")
    (case / "contracts/agreement.txt").write_text("contract", encoding="utf-8")
    authorize_test_case(case, purge_after_utc="2026-01-01T00:00:00+00:00",
        permitted_retained_paths=["outputs/final.pdf", "contracts/agreement.txt"])
    policy = configure_retention(case, {
        "schema_version": RETENTION_SCHEMA,
        "configured": True,
        "purge_after_utc": "2026-01-01T00:00:00+00:00",
        "mission_closed": True,
        "derived_retention_authorized": False,
        "retained_paths": ["outputs/final.pdf", "contracts/agreement.txt"],
    })
    assert policy["derived_retention_authorized"] is False
    receipt = purge_client_case(case, now=datetime(2026, 9, 5, tzinfo=timezone.utc))
    assert receipt["status"] == "complete"
    assert not (case / "sanitized/energy.csv").exists()
    assert not (case / "processed/cache.json").exists()
    assert not (case / "scratch/temp.txt").exists()
    assert (case / "outputs/final.pdf").is_file()
    assert (case / "contracts/agreement.txt").is_file()
    serialized = json.dumps(receipt)
    assert "client temporary" not in serialized and '"client": "value"' not in serialized
    assert json.loads((case / "PURGE_RECEIPT.json").read_text())["contains_client_content"] is False
    assert inspect_privacy_status(case)["state"] == "PURGED"
    with pytest.raises(ValueError, match="PURGED"):
        prepare_investigation(case / "outputs/final.pdf", case / "processed")
    new_drop = tmp_path / "new_drop"
    new_drop.mkdir()
    (new_drop / "new.csv").write_text("timestamp,power_kw\nT1,12\n", encoding="utf-8")
    with pytest.raises(ValueError, match="PURGED"):
        stage_incoming_drop(new_drop, case)


def test_purge_default_denies_derived_retention_and_reports_failures_without_paths(tmp_path: Path) -> None:
    case = _case(tmp_path)
    (case / "incoming/energy.csv").write_text("timestamp,power_kw\nT1,10\n", encoding="utf-8")
    validate_codex_privacy_review(case, _review(case, [_file_spec("incoming/energy.csv")], status="PASS"))
    (case / "retained_derived/derived.csv").write_text("site,volume\nUniqueSite,12345\n", encoding="utf-8")
    (case / "processed/private-cache.json").write_text('{"operator": "Jean Dupont"}', encoding="utf-8")
    authorize_test_case(case, purge_after_utc="2026-01-01T00:00:00+00:00")
    configure_retention(case, {
        "schema_version": RETENTION_SCHEMA,
        "configured": True,
        "purge_after_utc": "2026-01-01T00:00:00+00:00",
        "mission_closed": True,
        "derived_retention_authorized": False,
        "retained_paths": [],
    })

    def fail_one(path: Path) -> None:
        if path.name == "private-cache.json":
            raise PermissionError("simulated")
        path.unlink()

    receipt = purge_client_case(
        case, now=datetime(2026, 9, 5, tzinfo=timezone.utc), unlink_file=fail_one
    )
    serialized = json.dumps(receipt, ensure_ascii=False)
    assert receipt["status"] == "partial_failure"
    assert not (case / "retained_derived/derived.csv").exists()
    assert (case / "processed/private-cache.json").exists()
    assert "private-cache.json" not in serialized
    assert "Jean Dupont" not in serialized
    assert inspect_privacy_status(case)["state"] == "PRIVACY_BLOCKED"


def test_finding_generation_is_blocked_before_privacy_clearance(tmp_path: Path) -> None:
    case = tmp_path / "cases/client"
    create_client_case("client", root=tmp_path / "cases")
    with pytest.raises(ValueError, match="PRIVACY GATE"):
        record_structured_findings(case, [], no_finding={"reason": "insufficient_data"})


def test_resume_attribution_and_human_review_contracts_survive_privacy_gate(tmp_path: Path) -> None:
    case = _case(tmp_path)
    (case / "incoming/energy.csv").write_text("timestamp,energy_kwh\n2026-01-01,10\n2026-01-02,11\n", encoding="utf-8")
    validate_codex_privacy_review(case, _review(case, [_file_spec("incoming/energy.csv")], status="PASS"))
    prepare_investigation(case / "sanitized/energy.csv", case / "processed")
    record_existing_data_exhaustion(case, analysis_inventory_ref="prepared_analysis.json", reviewed_sources=["evidence_card.json"])
    component = AnonymousElectricalComponent(
        "C1", 10, (9, 11), 6, 18, None, None, None, (0, 1, 2, 3, 4),
        None, None, None, .9, .9, 8, "dependent", {}, ("signal.json",),
    )
    assets = [EquipmentRecord("PRESS_07", family="press", nominal_power_kw=10)]
    result = run_minimal_attribution(case, component=component, inventory=assets)
    assert result["status"] == "assessed"
    human = json.loads((case / "processed/human_review.json").read_text())
    assert human["approved_for_delivery"] is False
    assert_workflow_action_allowed(case, "asset_attribution")


def test_gitignore_explicitly_covers_all_client_privacy_zones() -> None:
    text = Path(".gitignore").read_text(encoding="utf-8")
    for fragment in (
        "client_cases/", "client_cases/*/incoming/", "client_cases/*/privacy/",
        "client_cases/*/sanitized/", "workspaces/*", "workspaces/*/incoming/",
        "workspaces/*/privacy/", "workspaces/*/sanitized/", "workspaces/*/processed/",
        "workspaces/*/scratch/", "workspaces/*/outputs/",
    ):
        assert fragment in text
    for path in (
        "client_cases/acme/incoming/private.csv",
        "client_cases/acme/privacy/review.json",
        "client_cases/acme/sanitized/approved.csv",
        "workspaces/acme/processed/evidence.json",
        "workspaces/acme/outputs/report.pdf",
    ):
        result = subprocess.run(
            ["git", "check-ignore", "--no-index", "-q", path], check=False
        )
        assert result.returncode == 0, path

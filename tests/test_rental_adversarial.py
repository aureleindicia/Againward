from copy import deepcopy
from datetime import datetime
from decimal import localcontext, ROUND_DOWN
import json

import pytest

from againward.core.workflow import fingerprint
from againward.domains.rental.models import RentalCase
from againward.domains.rental.pricing import build_expected_ledger
from againward.domains.rental.reconciliation import reconcile
from againward.domains.rental.ingestion import load_rental_case, build_evidence_dataset, EXTRACTION_SCHEMA
from benchmarking.rental_cases import rental_packet, return_evidence
from tests.test_rental_ingestion import write_packet


def test_decimal_results_do_not_depend_on_callers_context():
    packet = rental_packet(); return_evidence(packet)
    packet["terms"][0]["rate"] = "723.69"
    case = RentalCase.from_dict(packet)
    expected, evidence = reconcile(case), build_evidence_dataset(case).to_dict()
    with localcontext() as context:
        context.prec = 3; context.rounding = ROUND_DOWN
        assert reconcile(case) == expected
        assert build_evidence_dataset(case).to_dict() == evidence


def test_tiers_follow_documented_return_duration():
    packet = rental_packet(); return_evidence(packet)
    packet["terms"][0].update(tier_min_days=7)
    packet["terms"].append({**packet["terms"][0], "term_id": "SHORT", "tier_min_days": 0, "tier_max_days": 6,
                            "billing_unit": "DAY", "rate": "120"})
    expected = build_expected_ledger(RentalCase.from_dict(packet))["entries"][0]
    assert expected["term_ids"] == ["SHORT"]
    assert expected["amount"] == "480.00"


@pytest.mark.parametrize("mutation", ["invoice_return", "invoice_period", "unaccepted_period"])
def test_invoice_or_unaccepted_source_cannot_establish_operational_or_contract_authority(mutation):
    packet = rental_packet(); return_evidence(packet)
    if mutation == "invoice_return": packet["events"][0]["evidence_refs"] = packet["actual_charges"][0]["evidence_refs"]
    if mutation == "invoice_period": packet["periods"][0]["evidence_refs"] = packet["actual_charges"][0]["evidence_refs"]
    if mutation == "unaccepted_period": packet["documents"][0]["status"] = "PROPOSED"
    expected = build_expected_ledger(RentalCase.from_dict(packet))["entries"][0]
    assert expected["amount"] is None


def test_percentage_cycle_and_deep_dependencies_fail_with_specific_gap():
    packet = rental_packet(); base = packet["terms"][0]
    packet["terms"] += [{**base, "term_id": "F" + str(n), "charge_key": "f" + str(n), "charge_type": "SURCHARGE",
        "billing_unit": "PERCENT", "rate": "1", "percentage_of": "hire" if n == 0 else "f" + str(n - 1)} for n in range(40)]
    entries = build_expected_ledger(RentalCase.from_dict(packet))["entries"]
    assert any("percentage_dependency_depth_exceeds_32" in e["limitations"] for e in entries)
    packet["terms"][1]["percentage_of"] = "f39"
    entries = build_expected_ledger(RentalCase.from_dict(packet))["entries"]
    assert sum("missing_or_cyclic_percentage_base" in e["limitations"] for e in entries) == 40


def test_xlsx_explicit_mapping_calendar_dates_and_formula_refusal(tmp_path):
    from openpyxl import Workbook, load_workbook
    source, packet = write_packet(tmp_path)
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Invoice"
    sheet.append(["Invoice", "Line", "Amount", "Start", "End"])
    sheet.append(["I1", "L1", 850.25, datetime(2026, 9, 1), datetime(2026, 9, 8)])
    path = tmp_path / "invoice.xlsx"; workbook.save(path); workbook.close()
    packet["documents"][1].update(path=path.name, sha256=fingerprint(path))
    packet["actual_charges"] = []
    extraction = {"schema_version": EXTRACTION_SCHEMA, "case": packet, "tables": [{"document_id": "INVOICE", "record_type": "actual_charges",
        "sheet": "Invoice", "columns": {"invoice_id": "Invoice", "invoice_line_id": "Line", "net_amount": "Amount", "start": "Start", "end": "End"},
        "constants": {"period_id": "PERIOD", "charge_key": "hire", "charge_type": "RENTAL", "currency": "EUR"}}]}
    source.write_text(json.dumps(extraction))
    case, inventory = load_rental_case(source)
    assert case.actual_charges[0].start == "2026-09-01"
    assert case.actual_charges[0].net_amount == "850.25"
    assert case.actual_charges[0].evidence_refs[0].location == "sheet:Invoice/row:2"
    workbook = load_workbook(path); workbook["Invoice"]["C2"] = "=700+150.25"; workbook.save(path); workbook.close()
    packet["documents"][1]["sha256"] = fingerprint(path); source.write_text(json.dumps(extraction))
    with pytest.raises(ValueError, match="Formula"): load_rental_case(source)
